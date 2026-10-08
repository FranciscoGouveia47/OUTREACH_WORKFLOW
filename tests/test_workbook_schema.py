from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from outreach.workbook_schema import ensure_sheet_columns


class WorkbookSchemaTests(unittest.TestCase):
    def test_ensure_sheet_columns_adds_missing_headers(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "workbook.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Data"
            sheet.append(["Existing"])
            sheet.append(["value"])
            workbook.save(path)
            workbook.close()

            result = ensure_sheet_columns(path, "Data", ["Existing", "Added"])

            self.assertTrue(result.changed)
            self.assertEqual(("Added",), result.added_columns)
            updated = load_workbook(path, data_only=True)
            try:
                sheet = updated["Data"]
                self.assertEqual("Existing", sheet["A1"].value)
                self.assertEqual("Added", sheet["B1"].value)
                self.assertEqual("value", sheet["A2"].value)
            finally:
                updated.close()

    def test_ensure_sheet_columns_is_idempotent(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "workbook.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Data"
            sheet.append(["Existing", "Added"])
            workbook.save(path)
            workbook.close()

            result = ensure_sheet_columns(path, "Data", ["Existing", "Added"])

            self.assertFalse(result.changed)
            self.assertEqual((), result.added_columns)


if __name__ == "__main__":
    unittest.main()

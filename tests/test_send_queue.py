from __future__ import annotations

import csv
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from outreach.logger import (
    DELIVERY_FAILED_STATUS,
    DRAFT_CREATED_STATUS,
    OUTLOOK_SCHEDULED_STATUS,
    OUTLOOK_SCHEDULE_CANCELLED_STATUS,
    REFUSED_TO_ANSWER_STATUS,
    SENT_STATUS,
    UNKNOWN_STATUS,
    append_send_log,
)
from outreach.models import QueueEmail
from outreach.review_queue import (
    read_ready_queue,
    set_ready_to_send,
    set_ready_to_send_many,
    set_scheduled_send_at,
)
from outreach.scheduler import process_send_queue


def _queue_email(test_id: str = "TEST-001") -> QueueEmail:
    return QueueEmail(
        test_id=test_id,
        target_name="Target",
        recipient_email="test@example.com",
        recipient_type="Individual",
        template_type="General / Fallback",
        language="English",
        subject="[DRAFT] Subject",
        body="Body",
        ready_to_send=True,
        validation_status="VALID",
    )


class SendQueueTests(unittest.TestCase):
    def test_read_ready_queue_selects_only_ready_valid_rows(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "review.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Email Review Queue"
            sheet.append(
                [
                    "Test ID",
                    "Target Name",
                    "Recipient Email",
                    "Recipient Type",
                    "Template Type",
                    "Language",
                    "Subject",
                    "Email Body",
                    "Ready To Send",
                    "Validation Status",
                ]
            )
            sheet.append(
                [
                    "TEST-001",
                    "Target",
                    "test@example.com",
                    "Individual",
                    "General / Fallback",
                    "English",
                    "[DRAFT] Subject",
                    "Body",
                    "Yes",
                    "VALID",
                ]
            )
            sheet.append(
                [
                    "TEST-002",
                    "Target",
                    "blocked@example.com",
                    "Individual",
                    "General / Fallback",
                    "English",
                    "[DRAFT] Subject",
                    "Body",
                    "No",
                    "VALID",
                ]
            )
            workbook.save(path)

            ready = read_ready_queue(path)

            self.assertEqual(1, len(ready))
            self.assertEqual("TEST-001", ready[0].test_id)

    def test_set_ready_to_send_updates_review_queue(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "review.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Email Review Queue"
            sheet.append(["Test ID", "Ready To Send"])
            sheet.append(["TEST-001", "No"])
            workbook.save(path)
            workbook.close()

            set_ready_to_send(path, "TEST-001", True)

            updated = load_workbook(path, data_only=True)
            try:
                self.assertEqual("Yes", updated["Email Review Queue"]["B2"].value)
            finally:
                updated.close()

    def test_set_ready_to_send_keeps_original_when_atomic_save_fails(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "review.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Email Review Queue"
            sheet.append(["Test ID", "Ready To Send"])
            sheet.append(["TEST-001", "No"])
            workbook.save(path)
            workbook.close()

            with patch("outreach.workbook_io.load_workbook", side_effect=RuntimeError("validation failed")):
                with self.assertRaisesRegex(RuntimeError, "validation failed"):
                    set_ready_to_send(path, "TEST-001", True)

            updated = load_workbook(path, data_only=True)
            try:
                self.assertEqual("No", updated["Email Review Queue"]["B2"].value)
            finally:
                updated.close()
            self.assertEqual([], list(Path(temp_dir).glob(".*.tmp.xlsx")))

    def test_set_ready_to_send_many_updates_rows_with_one_save(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "review.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Email Review Queue"
            sheet.append(["Test ID", "Ready To Send"])
            sheet.append(["TEST-001", "No"])
            sheet.append(["TEST-002", "No"])
            workbook.save(path)
            workbook.close()

            with patch("outreach.review_queue.save_workbook_atomically") as save:
                set_ready_to_send_many(path, {"TEST-001": True, "TEST-002": True})

            save.assert_called_once()
            saved_workbook = save.call_args.args[0]
            try:
                sheet = saved_workbook["Email Review Queue"]
                self.assertEqual("Yes", sheet["B2"].value)
                self.assertEqual("Yes", sheet["B3"].value)
            finally:
                saved_workbook.close()

    def test_future_scheduled_row_is_not_ready_until_due(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "review.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Email Review Queue"
            sheet.append(
                [
                    "Test ID",
                    "Target Name",
                    "Recipient Email",
                    "Recipient Type",
                    "Template Type",
                    "Language",
                    "Subject",
                    "Email Body",
                    "Ready To Send",
                    "Scheduled Send At",
                    "Validation Status",
                ]
            )
            sheet.append(
                [
                    "TEST-001",
                    "Target",
                    "test@example.com",
                    "Individual",
                    "General / Fallback",
                    "English",
                    "[DRAFT] Subject",
                    "Body",
                    "Yes",
                    "2999-01-01 09:00",
                    "VALID",
                ]
            )
            workbook.save(path)
            workbook.close()

            self.assertEqual([], read_ready_queue(path))

    def test_past_scheduled_row_is_ready(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "review.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Email Review Queue"
            sheet.append(
                [
                    "Test ID",
                    "Target Name",
                    "Recipient Email",
                    "Recipient Type",
                    "Template Type",
                    "Language",
                    "Subject",
                    "Email Body",
                    "Ready To Send",
                    "Scheduled Send At",
                    "Validation Status",
                ]
            )
            sheet.append(
                [
                    "TEST-001",
                    "Target",
                    "test@example.com",
                    "Individual",
                    "General / Fallback",
                    "English",
                    "[DRAFT] Subject",
                    "Body",
                    "Yes",
                    "2000-01-01 09:00",
                    "VALID",
                ]
            )
            workbook.save(path)
            workbook.close()

            ready = read_ready_queue(path)
            self.assertEqual(1, len(ready))
            self.assertEqual("2000-01-01 09:00", ready[0].scheduled_send_at)

    def test_literal_none_scheduled_row_is_treated_as_unscheduled(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "review.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Email Review Queue"
            sheet.append(
                [
                    "Test ID",
                    "Target Name",
                    "Recipient Email",
                    "Recipient Type",
                    "Template Type",
                    "Language",
                    "Subject",
                    "Email Body",
                    "Ready To Send",
                    "Scheduled Send At",
                    "Validation Status",
                ]
            )
            sheet.append(
                [
                    "TEST-001",
                    "Target",
                    "test@example.com",
                    "Individual",
                    "General / Fallback",
                    "English",
                    "[DRAFT] Subject",
                    "Body",
                    "Yes",
                    "None",
                    "VALID",
                ]
            )
            workbook.save(path)
            workbook.close()

            ready = read_ready_queue(path)

            self.assertEqual(1, len(ready))
            self.assertEqual("", ready[0].scheduled_send_at)

    def test_set_scheduled_send_at_adds_column(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "review.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Email Review Queue"
            sheet.append(["Test ID"])
            sheet.append(["TEST-001"])
            workbook.save(path)
            workbook.close()

            set_scheduled_send_at(path, "TEST-001", "2026-08-18 09:30")

            updated = __import__("openpyxl", fromlist=["load_workbook"]).load_workbook(path, data_only=True)
            try:
                sheet = updated["Email Review Queue"]
                self.assertEqual("Scheduled Send At", sheet["B1"].value)
                self.assertEqual("2026-08-18 09:30", sheet["B2"].value)
            finally:
                updated.close()

    def test_dry_run_logs_without_graph_credentials_or_allowlist(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            log = Path(temp_dir) / "send_log.csv"

            with patch("config.SEND_LOG_PATH", log):
                planned, skipped, blocked = process_send_queue(
                    [_queue_email()],
                    confirm_send=False,
                    interval_minutes=0,
                    max_daily=30,
                    respect_window=False,
                )

            self.assertEqual((1, 0, 0), (planned, skipped, blocked))
            with log.open("r", encoding="utf-8", newline="") as handle:
                rows = list(csv.DictReader(handle))
            self.assertEqual("DRY_RUN", rows[0]["status"])

    def test_already_sent_rows_are_skipped(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            log = Path(temp_dir) / "send_log.csv"
            email = _queue_email()
            append_send_log(log, email, SENT_STATUS)

            with patch("config.SEND_LOG_PATH", log):
                planned, skipped, blocked = process_send_queue(
                    [email],
                    confirm_send=False,
                    interval_minutes=0,
                    max_daily=30,
                    respect_window=False,
                )

            self.assertEqual((0, 1, 0), (planned, skipped, blocked))

    def test_unknown_send_status_blocks_retry(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            log = Path(temp_dir) / "send_log.csv"
            email = _queue_email()
            append_send_log(log, email, UNKNOWN_STATUS)

            with patch("config.SEND_LOG_PATH", log):
                planned, skipped, blocked = process_send_queue(
                    [email],
                    confirm_send=False,
                    interval_minutes=0,
                    max_daily=30,
                    respect_window=False,
                )

            self.assertEqual((0, 0, 1), (planned, skipped, blocked))

    def test_delivery_failed_status_blocks_retry(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            log = Path(temp_dir) / "send_log.csv"
            email = _queue_email()
            append_send_log(log, email, DELIVERY_FAILED_STATUS)

            with patch("config.SEND_LOG_PATH", log):
                planned, skipped, blocked = process_send_queue(
                    [email],
                    confirm_send=False,
                    interval_minutes=0,
                    max_daily=30,
                    respect_window=False,
                )

            self.assertEqual((0, 0, 1), (planned, skipped, blocked))

    def test_stale_prompt_version_blocks_delivery(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            log = Path(temp_dir) / "send_log.csv"
            email = QueueEmail(**{**_queue_email().__dict__, "prompt_version": "old-version"})

            with (
                patch("config.SEND_LOG_PATH", log),
                patch("config.PROMPT_PACKET_VERSION", "current-version"),
            ):
                planned, skipped, blocked = process_send_queue(
                    [email],
                    confirm_send=False,
                    interval_minutes=0,
                    max_daily=30,
                    respect_window=False,
                )

            self.assertEqual((0, 0, 1), (planned, skipped, blocked))
            with log.open("r", encoding="utf-8", newline="") as handle:
                rows = list(csv.DictReader(handle))
            self.assertEqual("BLOCKED", rows[-1]["status"])
            self.assertIn("Stale prompt version old-version", rows[-1]["error"])

    def test_refused_to_answer_status_blocks_retry(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            log = Path(temp_dir) / "send_log.csv"
            email = _queue_email()
            append_send_log(log, email, REFUSED_TO_ANSWER_STATUS)

            with patch("config.SEND_LOG_PATH", log):
                planned, skipped, blocked = process_send_queue(
                    [email],
                    confirm_send=False,
                    interval_minutes=0,
                    max_daily=30,
                    respect_window=False,
                )

            self.assertEqual((0, 0, 1), (planned, skipped, blocked))

    def test_outlook_draft_mode_uses_outlook_client_and_logs_draft(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            log = Path(temp_dir) / "send_log.csv"

            with (
                patch("config.SEND_LOG_PATH", log),
                patch("outreach.scheduler.run_outlook_action") as outlook_action,
            ):
                outlook_action.return_value = "draft saved"
                planned, skipped, blocked = process_send_queue(
                    [_queue_email()],
                    confirm_send=True,
                    delivery="outlook-drafts",
                    interval_minutes=0,
                    max_daily=30,
                    respect_window=False,
                )

            self.assertEqual((1, 0, 0), (planned, skipped, blocked))
            with log.open("r", encoding="utf-8", newline="") as handle:
                rows = list(csv.DictReader(handle))
            self.assertEqual("DRAFT_CREATED", rows[0]["status"])

    def test_outlook_send_uses_logged_draft_entry_id(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            log = Path(temp_dir) / "send_log.csv"
            email = _queue_email()
            append_send_log(
                log,
                email,
                DRAFT_CREATED_STATUS,
                graph_response="Outlook draft saved (draft-entry-id-123)",
            )

            with (
                patch("config.SEND_LOG_PATH", log),
                patch("outreach.scheduler.run_outlook_action") as outlook_action,
            ):
                outlook_action.return_value = "fresh message sent; draft deleted"
                planned, skipped, blocked = process_send_queue(
                    [email],
                    confirm_send=True,
                    delivery="outlook",
                    interval_minutes=0,
                    max_daily=30,
                    respect_window=False,
                )

            self.assertEqual((1, 0, 0), (planned, skipped, blocked))
            outlook_action.assert_called_once_with(
                email,
                "send-draft",
                60,
                draft_entry_id="draft-entry-id-123",
            )

    def test_outlook_scheduled_mode_uses_logged_draft_entry_id_and_logs_scheduled(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            log = Path(temp_dir) / "send_log.csv"
            email = _queue_email()
            email = QueueEmail(**{**email.__dict__, "scheduled_send_at": "2999-01-01 09:00"})
            append_send_log(
                log,
                email,
                DRAFT_CREATED_STATUS,
                graph_response="Outlook draft saved (draft-entry-id-123)",
            )

            with (
                patch("config.SEND_LOG_PATH", log),
                patch("outreach.scheduler.run_outlook_action") as outlook_action,
            ):
                outlook_action.return_value = "fresh deferred message requested; draft deleted"
                planned, skipped, blocked = process_send_queue(
                    [email],
                    confirm_send=True,
                    delivery="outlook-scheduled",
                    interval_minutes=0,
                    max_daily=30,
                    respect_window=False,
                )

            self.assertEqual((1, 0, 0), (planned, skipped, blocked))
            outlook_action.assert_called_once_with(
                email,
                "schedule-draft",
                60,
                draft_entry_id="draft-entry-id-123",
            )
            with log.open("r", encoding="utf-8", newline="") as handle:
                rows = list(csv.DictReader(handle))
            self.assertEqual(OUTLOOK_SCHEDULED_STATUS, rows[-1]["status"])
            self.assertEqual("2999-01-01 09:00", rows[-1]["scheduled_send_at"])

    def test_outlook_scheduled_mode_blocks_blank_schedule(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            log = Path(temp_dir) / "send_log.csv"

            with (
                patch("config.SEND_LOG_PATH", log),
                patch("outreach.scheduler.run_outlook_action") as outlook_action,
            ):
                planned, skipped, blocked = process_send_queue(
                    [_queue_email()],
                    confirm_send=True,
                    delivery="outlook-scheduled",
                    interval_minutes=0,
                    max_daily=30,
                    respect_window=False,
                )

            self.assertEqual((0, 0, 1), (planned, skipped, blocked))
            outlook_action.assert_not_called()
            with log.open("r", encoding="utf-8", newline="") as handle:
                rows = list(csv.DictReader(handle))
            self.assertEqual("BLOCKED", rows[-1]["status"])

    def test_append_send_log_writes_blank_scheduled_time_when_value_is_none(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            log = Path(temp_dir) / "send_log.csv"
            email = QueueEmail(**{**_queue_email().__dict__, "scheduled_send_at": None})

            append_send_log(log, email, DRAFT_CREATED_STATUS)

            with log.open("r", encoding="utf-8", newline="") as handle:
                rows = list(csv.DictReader(handle))
            self.assertEqual("", rows[0]["scheduled_send_at"])

    def test_cancelled_outlook_schedule_no_longer_skips_retry(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            log = Path(temp_dir) / "send_log.csv"
            email = _queue_email()
            append_send_log(log, email, OUTLOOK_SCHEDULED_STATUS)
            append_send_log(log, email, OUTLOOK_SCHEDULE_CANCELLED_STATUS)

            with patch("config.SEND_LOG_PATH", log):
                planned, skipped, blocked = process_send_queue(
                    [email],
                    confirm_send=False,
                    interval_minutes=0,
                    max_daily=30,
                    respect_window=False,
                )

            self.assertEqual((1, 0, 0), (planned, skipped, blocked))

    def test_cancelled_outlook_schedule_allows_new_draft(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            log = Path(temp_dir) / "send_log.csv"
            email = _queue_email()
            append_send_log(
                log,
                email,
                DRAFT_CREATED_STATUS,
                graph_response="Outlook draft saved (old-draft-entry-id)",
            )
            append_send_log(log, email, OUTLOOK_SCHEDULED_STATUS)
            append_send_log(log, email, OUTLOOK_SCHEDULE_CANCELLED_STATUS)

            with (
                patch("config.SEND_LOG_PATH", log),
                patch("outreach.scheduler.run_outlook_action") as outlook_action,
            ):
                outlook_action.return_value = "Outlook draft saved (new-draft-entry-id)"
                planned, skipped, blocked = process_send_queue(
                    [email],
                    confirm_send=True,
                    delivery="outlook-drafts",
                    interval_minutes=0,
                    max_daily=30,
                    respect_window=False,
                )

            self.assertEqual((1, 0, 0), (planned, skipped, blocked))
            outlook_action.assert_called_once_with(email, "draft", 60)

    def test_schedule_after_cancel_uses_new_draft_entry_id(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            log = Path(temp_dir) / "send_log.csv"
            email = QueueEmail(**{**_queue_email().__dict__, "scheduled_send_at": "2999-01-01 09:00"})
            append_send_log(
                log,
                email,
                DRAFT_CREATED_STATUS,
                graph_response="Outlook draft saved (old-draft-entry-id)",
            )
            append_send_log(log, email, OUTLOOK_SCHEDULED_STATUS)
            append_send_log(log, email, OUTLOOK_SCHEDULE_CANCELLED_STATUS)
            append_send_log(
                log,
                email,
                DRAFT_CREATED_STATUS,
                graph_response="Outlook draft saved (new-draft-entry-id)",
            )

            with (
                patch("config.SEND_LOG_PATH", log),
                patch("outreach.scheduler.run_outlook_action") as outlook_action,
            ):
                outlook_action.return_value = "fresh deferred message requested; draft deleted"
                planned, skipped, blocked = process_send_queue(
                    [email],
                    confirm_send=True,
                    delivery="outlook-scheduled",
                    interval_minutes=0,
                    max_daily=30,
                    respect_window=False,
                )

            self.assertEqual((1, 0, 0), (planned, skipped, blocked))
            outlook_action.assert_called_once_with(
                email,
                "schedule-draft",
                60,
                draft_entry_id="new-draft-entry-id",
            )

    def test_append_send_log_migrates_legacy_log_and_derives_scheduled_time(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            log = Path(temp_dir) / "send_log.csv"
            log.write_text(
                "timestamp,test_id,target_name,recipient_email,recipient_type,template_type,language,subject,status,graph_response,error\n"
                "2000-01-01T08:55:00+00:00,TEST-000,Target,test@example.com,Individual,General / Fallback,English,Subject,OUTLOOK_SCHEDULED,Outlook deferred send requested (2000-01-01 09:00),\n",
                encoding="utf-8",
            )
            email = QueueEmail(
                **{
                    **_queue_email("TEST-001").__dict__,
                    "scheduled_send_at": "2999-01-01 09:00",
                }
            )

            append_send_log(log, email, OUTLOOK_SCHEDULED_STATUS, graph_response="scheduled")

            with log.open("r", encoding="utf-8", newline="") as handle:
                rows = list(csv.DictReader(handle))
            self.assertIn("scheduled_send_at", rows[0])
            self.assertEqual("2000-01-01 09:00", rows[0]["scheduled_send_at"])
            self.assertEqual("2999-01-01 09:00", rows[1]["scheduled_send_at"])


if __name__ == "__main__":
    unittest.main()

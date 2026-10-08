from __future__ import annotations

import tempfile
import unittest
from dataclasses import replace
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from main import _generate_handoff
from outreach.excel_reader import read_workbook
from outreach.handoff_workbook import (
    INSTRUCTIONS_SHEET_NAME,
    SHEET_NAME,
    remove_handoff_email,
    upsert_handoff_workbook_rows,
    write_handoff_workbook,
)
from outreach.lineage import handoff_id_from_workbook, read_metadata, write_metadata
from outreach.logger import (
    DELIVERY_FAILED_STATUS,
    OUTLOOK_SCHEDULED_STATUS,
    REFUSED_TO_ANSWER_STATUS,
    SENT_STATUS,
    append_send_log,
)
from outreach.models import HandoffEmail, QueueEmail
from outreach.prompt_packet import build_prompt_packet
from outreach.review_importer import import_completed_handoff
from outreach.template_engine import render_email

import config


class HandoffWorkflowTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.subjects, cls.templates = read_workbook(config.EXCEL_PATH)
        cls.subject = next(item for item in cls.subjects if item.enabled)
        cls.email = render_email(cls.subject, cls.templates)
        cls.handoff = HandoffEmail(
            stakeholder=cls.subject,
            email=cls.email,
            prompt_packet=build_prompt_packet(cls.subject, cls.email),
        )

    def test_handoff_workbook_excludes_real_recipient_email(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "handoff.xlsx"

            write_handoff_workbook([self.handoff], path, datetime(2026, 8, 17, 14, 0, 0))

            workbook = load_workbook(path, data_only=True)
            sheet = workbook[SHEET_NAME]
            headers = [sheet.cell(1, column).value for column in range(1, sheet.max_column + 1)]
            row_values = [
                str(sheet.cell(2, column).value or "")
                for column in range(1, sheet.max_column + 1)
            ]
            self.assertNotIn("Recipient Email", headers)
            self.assertNotIn(self.subject.recipient_email, "\n".join(row_values))
            self.assertIn("ChatGPT Prompt Packet", headers)
            self.assertIn("AI Risk Flags", headers)
            self.assertIn("Prompt Version", headers)

    def test_handoff_workbook_contains_general_instructions_sheet(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "handoff.xlsx"

            write_handoff_workbook([self.handoff], path, datetime(2026, 8, 17, 14, 0, 0))

            workbook = load_workbook(path, data_only=True)
            self.assertIn(INSTRUCTIONS_SHEET_NAME, workbook.sheetnames)
            instructions = workbook[INSTRUCTIONS_SHEET_NAME]["A2"].value
            normalised_instructions = " ".join(instructions.split())
            self.assertIn("Fill only these columns", instructions)
            self.assertIn("AI Subject", instructions)
            self.assertIn("AI Risk Flags", instructions)
            self.assertIn("Do not mechanically paste the Organisation field", instructions)
            self.assertIn("only source available in this handoff", instructions)
            self.assertIn("Contact-fit and personalisation guardrails", instructions)
            self.assertIn("perform a local relevance check", instructions)
            self.assertIn("weak individual evidence", normalised_instructions)
            self.assertIn("referral request", instructions)
            self.assertIn("Never create personalisation by concatenating workbook fields", instructions)
            self.assertIn("beyond what is already in the workbook row", instructions)
            self.assertIn("verify it externally", instructions)
            self.assertIn("capital letters that appear unnecessarily", instructions)
            self.assertIn("Use normal sentence case", instructions)
            self.assertIn("read like mechanically repeated bulk mail", instructions)
            self.assertIn("deceptive tactics", instructions)
            self.assertIn("supplier operations across EMEA", instructions)
            self.assertIn("supplier onboarding in a manufacturing context", instructions)
            self.assertIn("Do not describe the evidence source", instructions)
            self.assertIn("your LinkedIn says", instructions)
            self.assertIn("your work aligns with our solution", instructions)

    def test_prompt_packet_contains_campaign_locked_values(self) -> None:
        prompt = build_prompt_packet(self.subject, self.email)

        self.assertIn(f"Campaign name: {config.CAMPAIGN_NAME}", prompt)
        self.assertIn(f"Call-to-action URL: {config.CTA_URL}", prompt)
        self.assertIn(f"Company website: {config.COMPANY_URL}", prompt)

    def test_prompt_packet_contains_natural_affiliation_rule(self) -> None:
        prompt = build_prompt_packet(self.subject, self.email)
        normalised_prompt = " ".join(prompt.split())

        self.assertIn("Do not mechanically paste the Organisation field", prompt)
        self.assertIn("only source available in this handoff", prompt)
        self.assertIn("Contact-fit and personalisation guardrails", prompt)
        self.assertIn("perform a local relevance check", prompt)
        self.assertIn("weak individual evidence", normalised_prompt)
        self.assertIn("referral request", prompt)
        self.assertIn("Never create personalisation by concatenating workbook fields", prompt)
        self.assertIn("beyond what is already in the stakeholder context", prompt)
        self.assertIn("verify it externally", prompt)
        self.assertIn("capital letters that appear unnecessarily", prompt)
        self.assertIn("Use normal sentence case", prompt)
        self.assertIn("do not read like mechanically repeated bulk mail", prompt)
        self.assertIn("Do not use deceptive tactics", prompt)
        self.assertIn("Your work at Acme Group, Procurement / EMEA Supplier Operations", prompt)
        self.assertIn("Given your work on supplier operations across EMEA", prompt)
        self.assertIn("use the role, field, or personalisation fact instead", prompt)
        self.assertIn("Do not describe the evidence source", prompt)
        self.assertIn("your LinkedIn says", prompt)
        self.assertIn("your work aligns with our solution", prompt)

    def test_upsert_handoff_rows_replaces_selected_target_only(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "handoff.xlsx"
            second_subject = next(
                item
                for item in self.subjects
                if item.enabled and item.test_id != self.subject.test_id
            )
            second_email = render_email(second_subject, self.templates)
            second_handoff = HandoffEmail(
                stakeholder=second_subject,
                email=second_email,
                prompt_packet=build_prompt_packet(second_subject, second_email),
            )
            write_handoff_workbook([self.handoff, second_handoff], path, datetime(2026, 8, 17, 14, 0, 0))

            replacement = HandoffEmail(
                stakeholder=self.subject,
                email=self.email,
                prompt_packet="replacement prompt packet",
            )
            upsert_handoff_workbook_rows([replacement], path, datetime(2026, 8, 17, 15, 0, 0))

            workbook = load_workbook(path, data_only=True)
            sheet = workbook[SHEET_NAME]
            rows = {
                sheet.cell(row, 1).value: sheet.cell(row, 13).value
                for row in range(2, sheet.max_row + 1)
            }
            self.assertEqual(2, len(rows))
            self.assertEqual("replacement prompt packet", rows[self.subject.test_id])
            self.assertIn(second_subject.test_id, rows)

    def test_remove_handoff_email_removes_one_target_from_handoff_workbook(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "handoff.xlsx"
            write_handoff_workbook([self.handoff], path, datetime(2026, 8, 17, 14, 0, 0))

            removed = remove_handoff_email(path, self.subject.test_id)

            self.assertTrue(removed)
            updated = load_workbook(path, data_only=True)
            sheet = updated[SHEET_NAME]
            values = [sheet.cell(row, 1).value for row in range(2, sheet.max_row + 1)]
            self.assertNotIn(self.subject.test_id, values)

    def test_import_completed_handoff_creates_local_review_queue(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            handoff_path = Path(temp_dir) / "handoff.xlsx"
            review_path = Path(temp_dir) / "review.xlsx"
            write_handoff_workbook([self.handoff], handoff_path, datetime(2026, 8, 17, 14, 0, 0))

            workbook = load_workbook(handoff_path)
            sheet = workbook[SHEET_NAME]
            sheet["N2"] = self.email.subject
            sheet["O2"] = self.email.body
            sheet["P2"] = "Template-only test output."
            sheet["Q2"] = "Yes"
            workbook.save(handoff_path)

            import_completed_handoff(
                handoff_path,
                self.subjects,
                review_path,
                datetime(2026, 8, 17, 14, 30, 0),
            )

            review = load_workbook(review_path, data_only=True)
            sheet = review["Email Review Queue"]
            self.assertEqual(self.subject.recipient_email, sheet["C2"].value)
            self.assertEqual("No", sheet["I2"].value)
            self.assertEqual("VALID", sheet["M2"].value)
            self.assertEqual(config.PROMPT_PACKET_VERSION, sheet["R2"].value)

    def test_import_completed_handoff_blocks_stale_prompt_version(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            handoff_path = Path(temp_dir) / "handoff.xlsx"
            review_path = Path(temp_dir) / "review.xlsx"
            write_handoff_workbook([self.handoff], handoff_path, datetime(2026, 8, 17, 14, 0, 0))

            workbook = load_workbook(handoff_path)
            sheet = workbook[SHEET_NAME]
            sheet["N2"] = self.email.subject
            sheet["O2"] = self.email.body
            sheet["P2"] = "Template-only test output."
            sheet["Q2"] = "Yes"
            sheet["T2"] = "old-version"
            workbook.save(handoff_path)

            with patch("config.PROMPT_PACKET_VERSION", "current-version"):
                import_completed_handoff(
                    handoff_path,
                    self.subjects,
                    review_path,
                    datetime(2026, 8, 17, 14, 30, 0),
                )

            review = load_workbook(review_path, data_only=True)
            sheet = review["Email Review Queue"]
            self.assertEqual("INVALID", sheet["M2"].value)
            self.assertIn("Prompt Version is stale", sheet["N2"].value)

    def test_import_completed_handoff_skips_already_sent_rows(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            handoff_path = Path(temp_dir) / "handoff.xlsx"
            review_path = Path(temp_dir) / "review.xlsx"
            send_log_path = Path(temp_dir) / "send_log.csv"
            write_handoff_workbook([self.handoff], handoff_path, datetime(2026, 8, 17, 14, 0, 0))

            workbook = load_workbook(handoff_path)
            sheet = workbook[SHEET_NAME]
            sheet["N2"] = self.email.subject
            sheet["O2"] = self.email.body
            sheet["P2"] = "Template-only test output."
            sheet["Q2"] = "Yes"
            workbook.save(handoff_path)

            append_send_log(
                send_log_path,
                QueueEmail(
                    test_id=self.subject.test_id,
                    target_name=self.subject.target_name,
                    recipient_email=self.subject.recipient_email,
                    recipient_type=self.subject.recipient_type,
                    template_type=self.subject.template_type,
                    language=self.subject.language,
                    subject=self.email.subject,
                    body=self.email.body,
                    ready_to_send=True,
                    validation_status="VALID",
                ),
                SENT_STATUS,
            )

            _review_path, review_emails = import_completed_handoff(
                handoff_path,
                self.subjects,
                review_path,
                datetime(2026, 8, 17, 14, 30, 0),
                send_log_path=send_log_path,
            )

            review = load_workbook(review_path, data_only=True)
            sheet = review["Email Review Queue"]
            self.assertEqual([], review_emails)
            self.assertEqual(1, sheet.max_row)

    def test_generate_handoff_skips_delivery_failed_rows(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            handoff_path = Path(temp_dir) / "handoff.xlsx"
            send_log_path = Path(temp_dir) / "send_log.csv"
            append_send_log(
                send_log_path,
                QueueEmail(
                    test_id=self.subject.test_id,
                    target_name=self.subject.target_name,
                    recipient_email=self.subject.recipient_email,
                    recipient_type=self.subject.recipient_type,
                    template_type=self.subject.template_type,
                    language=self.subject.language,
                    subject=self.email.subject,
                    body=self.email.body,
                    ready_to_send=True,
                    validation_status="VALID",
                ),
                DELIVERY_FAILED_STATUS,
            )

            with (
                patch("config.CHATGPT_HANDOFF_WORKBOOK_PATH", handoff_path),
                patch("config.SEND_LOG_PATH", send_log_path),
            ):
                exit_code = _generate_handoff(
                    [self.subject],
                    self.templates,
                    [],
                    include_generated=True,
                    limit=None,
                )

            self.assertEqual(0, exit_code)
            self.assertFalse(handoff_path.exists())

    def test_generate_handoff_skips_outlook_scheduled_rows(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            handoff_path = Path(temp_dir) / "handoff.xlsx"
            send_log_path = Path(temp_dir) / "send_log.csv"
            append_send_log(
                send_log_path,
                QueueEmail(
                    test_id=self.subject.test_id,
                    target_name=self.subject.target_name,
                    recipient_email=self.subject.recipient_email,
                    recipient_type=self.subject.recipient_type,
                    template_type=self.subject.template_type,
                    language=self.subject.language,
                    subject=self.email.subject,
                    body=self.email.body,
                    ready_to_send=True,
                    validation_status="VALID",
                    scheduled_send_at="2026-09-23 09:00",
                ),
                OUTLOOK_SCHEDULED_STATUS,
            )

            with (
                patch("config.CHATGPT_HANDOFF_WORKBOOK_PATH", handoff_path),
                patch("config.SEND_LOG_PATH", send_log_path),
            ):
                exit_code = _generate_handoff(
                    [self.subject],
                    self.templates,
                    [],
                    include_generated=True,
                    limit=None,
                )

            self.assertEqual(0, exit_code)
            self.assertFalse(handoff_path.exists())

    def test_generate_handoff_skips_refused_to_answer_rows(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            handoff_path = Path(temp_dir) / "handoff.xlsx"
            send_log_path = Path(temp_dir) / "send_log.csv"
            append_send_log(
                send_log_path,
                QueueEmail(
                    test_id=self.subject.test_id,
                    target_name=self.subject.target_name,
                    recipient_email=self.subject.recipient_email,
                    recipient_type=self.subject.recipient_type,
                    template_type=self.subject.template_type,
                    language=self.subject.language,
                    subject=self.email.subject,
                    body=self.email.body,
                    ready_to_send=True,
                    validation_status="VALID",
                ),
                REFUSED_TO_ANSWER_STATUS,
            )

            with (
                patch("config.CHATGPT_HANDOFF_WORKBOOK_PATH", handoff_path),
                patch("config.SEND_LOG_PATH", send_log_path),
            ):
                exit_code = _generate_handoff(
                    [self.subject],
                    self.templates,
                    [],
                    include_generated=True,
                    limit=None,
                )

            self.assertEqual(0, exit_code)
            self.assertFalse(handoff_path.exists())

    def test_import_completed_handoff_blocks_recipient_email_leak(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            handoff_path = Path(temp_dir) / "handoff.xlsx"
            review_path = Path(temp_dir) / "review.xlsx"
            write_handoff_workbook([self.handoff], handoff_path, datetime(2026, 8, 17, 14, 0, 0))

            workbook = load_workbook(handoff_path)
            sheet = workbook[SHEET_NAME]
            sheet["N2"] = self.email.subject
            sheet["O2"] = self.email.body + f"\nContact: {self.subject.recipient_email}"
            sheet["P2"] = "Template-only test output."
            sheet["Q2"] = "Yes"
            workbook.save(handoff_path)

            import_completed_handoff(
                handoff_path,
                self.subjects,
                review_path,
                datetime(2026, 8, 17, 14, 30, 0),
            )

            review = load_workbook(review_path, data_only=True)
            sheet = review["Email Review Queue"]
            self.assertEqual("INVALID", sheet["M2"].value)
            self.assertIn("recipient email", sheet["N2"].value)

    def test_import_completed_handoff_flags_casual_localised_greeting(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            handoff_path = Path(temp_dir) / "handoff.xlsx"
            review_path = Path(temp_dir) / "review.xlsx"
            org_subject = replace(
                self.subject,
                recipient_type="Organisation",
                organisation="3B's Research Group",
                language="Portuguese",
            )
            org_email = render_email(org_subject, self.templates)
            write_handoff_workbook(
                [
                    HandoffEmail(
                        stakeholder=org_subject,
                        email=org_email,
                        prompt_packet=build_prompt_packet(org_subject, org_email),
                    )
                ],
                handoff_path,
                datetime(2026, 8, 17, 14, 0, 0),
            )

            workbook = load_workbook(handoff_path)
            sheet = workbook[SHEET_NAME]
            sheet["N2"] = org_email.subject
            sheet["O2"] = (
                "Olá, equipa de 3B's Research Group,\n\n"
                "Podem partilhar este convite com a pessoa mais adequada da equipa.\n"
                f"Survey: {config.SURVEY_URL}\n"
                f"Website: {config.PROJECT_URL}\n"
                f"LinkedIn: {config.PROJECT_LINKEDIN_URL}\n"
                f"Grant Agreement No. {config.PROJECT_GRANT_AGREEMENT}\n\n"
                f"Com os melhores cumprimentos,\n{config.SENDER_NAME}"
            )
            sheet["P2"] = "Test output."
            sheet["Q2"] = "Yes"
            workbook.save(handoff_path)

            import_completed_handoff(
                handoff_path,
                [org_subject],
                review_path,
                datetime(2026, 8, 17, 14, 30, 0),
            )

            review = load_workbook(review_path, data_only=True)
            sheet = review["Email Review Queue"]
            self.assertEqual("INVALID", sheet["M2"].value)
            self.assertIn("too casual", sheet["N2"].value)

    def test_import_completed_handoff_flags_organisation_singular_wording(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            handoff_path = Path(temp_dir) / "handoff.xlsx"
            review_path = Path(temp_dir) / "review.xlsx"
            org_subject = replace(
                self.subject,
                recipient_type="Organisation",
                organisation="3B's Research Group",
                language="Portuguese",
            )
            org_email = render_email(org_subject, self.templates)
            write_handoff_workbook(
                [
                    HandoffEmail(
                        stakeholder=org_subject,
                        email=org_email,
                        prompt_packet=build_prompt_packet(org_subject, org_email),
                    )
                ],
                handoff_path,
                datetime(2026, 8, 17, 14, 0, 0),
            )

            workbook = load_workbook(handoff_path)
            sheet = workbook[SHEET_NAME]
            sheet["N2"] = org_email.subject
            sheet["O2"] = (
                "Cara equipa de 3B's Research Group,\n\n"
                "Podem partilhar este convite com a pessoa mais adequada da equipa.\n"
                "Se preferir não abrir diretamente a ligação, pode consultar a página do projeto.\n"
                f"Survey: {config.SURVEY_URL}\n"
                f"Website: {config.PROJECT_URL}\n"
                f"LinkedIn: {config.PROJECT_LINKEDIN_URL}\n"
                f"Grant Agreement No. {config.PROJECT_GRANT_AGREEMENT}\n\n"
                f"Com os melhores cumprimentos,\n{config.SENDER_NAME}"
            )
            sheet["P2"] = "Test output."
            sheet["Q2"] = "Yes"
            workbook.save(handoff_path)

            import_completed_handoff(
                handoff_path,
                [org_subject],
                review_path,
                datetime(2026, 8, 17, 14, 30, 0),
            )

            review = load_workbook(review_path, data_only=True)
            sheet = review["Email Review Queue"]
            self.assertEqual("INVALID", sheet["M2"].value)
            self.assertIn("singular-person", sheet["N2"].value)

    def test_import_completed_handoff_flags_source_artifact_wording(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            handoff_path = Path(temp_dir) / "handoff.xlsx"
            review_path = Path(temp_dir) / "review.xlsx"
            write_handoff_workbook([self.handoff], handoff_path, datetime(2026, 8, 17, 14, 0, 0))

            workbook = load_workbook(handoff_path)
            sheet = workbook[SHEET_NAME]
            sheet["N2"] = self.email.subject
            sheet["O2"] = (
                f"{self.email.body}\n\n"
                "The university directly lists his expertise in biomechanics and additive manufacturing. "
                "For that reason, your perspective would be valuable to the consultation."
            )
            sheet["P2"] = "Test output."
            sheet["Q2"] = "Yes"
            workbook.save(handoff_path)

            import_completed_handoff(
                handoff_path,
                self.subjects,
                review_path,
                datetime(2026, 8, 17, 14, 30, 0),
            )

            review = load_workbook(review_path, data_only=True)
            sheet = review["Email Review Queue"]
            self.assertEqual("INVALID", sheet["M2"].value)
            self.assertIn("source evidence", sheet["N2"].value)

    def test_import_completed_handoff_flags_generic_fit_rationale(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            handoff_path = Path(temp_dir) / "handoff.xlsx"
            review_path = Path(temp_dir) / "review.xlsx"
            write_handoff_workbook([self.handoff], handoff_path, datetime(2026, 8, 17, 14, 0, 0))

            workbook = load_workbook(handoff_path)
            sheet = workbook[SHEET_NAME]
            sheet["N2"] = self.email.subject
            sheet["O2"] = (
                f"{self.email.body}\n\n"
                "Your work in additive manufacturing aligns closely with the areas explored "
                "in the consultation, and your input would be valuable."
            )
            sheet["P2"] = "Test output."
            sheet["Q2"] = "Yes"
            workbook.save(handoff_path)

            import_completed_handoff(
                handoff_path,
                self.subjects,
                review_path,
                datetime(2026, 8, 17, 14, 30, 0),
            )

            review = load_workbook(review_path, data_only=True)
            sheet = review["Email Review Queue"]
            self.assertEqual("INVALID", sheet["M2"].value)
            self.assertIn("source evidence", sheet["N2"].value)

    def test_repo_handoff_can_be_fingerprinted(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "handoff.xlsx"
            write_handoff_workbook([self.handoff], path, datetime(2026, 8, 17, 14, 0, 0))
            handoff_id = handoff_id_from_workbook(path)

        self.assertEqual(64, len(handoff_id))

    def test_metadata_reader_does_not_require_reported_sheet_dimensions(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "metadata.xlsx"
            from openpyxl import Workbook

            workbook = Workbook()
            write_metadata(workbook, {"file_type": "chatgpt_handoff", "handoff_id": "abc123"})
            workbook.save(path)
            workbook.close()

            metadata = read_metadata(path)

            self.assertEqual("chatgpt_handoff", metadata["file_type"])
            self.assertEqual("abc123", metadata["handoff_id"])


if __name__ == "__main__":
    unittest.main()

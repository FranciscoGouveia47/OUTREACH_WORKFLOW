from __future__ import annotations

import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import MagicMock, Mock, patch

from openpyxl import Workbook, load_workbook

import config
from outreach.handoff_workbook import write_handoff_workbook
from outreach.logger import append_send_log
from outreach.models import HandoffEmail, QueueEmail
from outreach.prompt_packet import build_prompt_packet
from outreach.review_importer import LOCAL_REVIEW_HEADERS, LOCAL_REVIEW_SHEET_NAME
from outreach.status import _issue_summary, _next_action, _read_review_rows, load_target_statuses
from outreach.template_engine import render_email


class StatusTests(unittest.TestCase):
    def test_read_review_rows_closes_workbook(self) -> None:
        class FakeCell:
            def __init__(self, value: str) -> None:
                self.value = value

        class FakeSheet:
            max_column = 2
            max_row = 2

            def cell(self, row: int, column: int) -> FakeCell:
                values = {
                    (1, 1): "Test ID",
                    (1, 2): "Ready To Send",
                    (2, 1): "TEST-001",
                    (2, 2): "Yes",
                }
                return FakeCell(values.get((row, column), ""))

        fake_workbook = MagicMock()
        fake_workbook.sheetnames = ["Email Review Queue"]
        fake_workbook.__getitem__ = Mock(return_value=FakeSheet())

        with tempfile.TemporaryDirectory() as temp_dir:
            review_path = Path(temp_dir) / "review.xlsx"
            review_path.touch()
            with patch("openpyxl.load_workbook", return_value=fake_workbook):
                rows = _read_review_rows(review_path)

        self.assertEqual("Yes", rows["TEST-001"]["Ready To Send"])
        fake_workbook.close.assert_called_once_with()

    def test_status_joins_source_review_and_send_log(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.xlsx"
            handoff = Path(temp_dir) / "handoff.xlsx"
            review = Path(temp_dir) / "review.xlsx"
            log = Path(temp_dir) / "send_log.csv"

            workbook = load_workbook(config.EXCEL_PATH)
            workbook.save(source)
            subjects, templates = __import__("outreach.excel_reader", fromlist=["read_workbook"]).read_workbook(source)
            subject = next(item for item in subjects if item.enabled)
            email_message = render_email(subject, templates)
            write_handoff_workbook(
                [
                    HandoffEmail(
                        stakeholder=subject,
                        email=email_message,
                        prompt_packet=build_prompt_packet(subject, email_message),
                    )
                ],
                handoff,
                datetime(2026, 8, 17, 12, 0, 0),
            )

            review_book = Workbook()
            sheet = review_book.active
            sheet.title = LOCAL_REVIEW_SHEET_NAME
            sheet.append(LOCAL_REVIEW_HEADERS)
            sheet.append(
                [
                    subject.test_id,
                    subject.target_name,
                    subject.recipient_email,
                    subject.recipient_type,
                    subject.template_type,
                    subject.language,
                    "Subject",
                    "Body",
                    "Yes",
                    "",
                    "Yes",
                    "",
                    "VALID",
                    "",
                    "2026-08-17T12:05:00+01:00",
                    "",
                ]
            )
            review_book.save(review)

            email = QueueEmail(
                test_id=subject.test_id,
                target_name=subject.target_name,
                recipient_email=subject.recipient_email,
                recipient_type=subject.recipient_type,
                template_type=subject.template_type,
                language=subject.language,
                subject="Subject",
                body="Body",
                ready_to_send=True,
                validation_status="VALID",
            )
            append_send_log(log, email, "SENT", graph_response="test")

            statuses = load_target_statuses(source, handoff, review, log)
            status = next(item for item in statuses if item.test_id == subject.test_id)

            self.assertEqual("Yes", status.handoff_generated)
            self.assertEqual("Yes", status.review_imported)
            self.assertEqual("Yes", status.ready_to_send)
            self.assertEqual("VALID", status.validation_status)
            self.assertEqual("", status.scheduled_send_at)
            self.assertEqual("SENT", status.last_send_status)
            self.assertEqual("", status.issues)
            self.assertEqual("Done", status.next_action)

    def test_status_derives_scheduled_time_from_legacy_log_response(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.xlsx"
            handoff = Path(temp_dir) / "handoff.xlsx"
            review = Path(temp_dir) / "review.xlsx"
            log = Path(temp_dir) / "send_log.csv"

            workbook = load_workbook(config.EXCEL_PATH)
            workbook.save(source)
            subjects, templates = __import__("outreach.excel_reader", fromlist=["read_workbook"]).read_workbook(source)
            subject = next(item for item in subjects if item.enabled)
            email_message = render_email(subject, templates)
            write_handoff_workbook(
                [
                    HandoffEmail(
                        stakeholder=subject,
                        email=email_message,
                        prompt_packet=build_prompt_packet(subject, email_message),
                    )
                ],
                handoff,
                datetime(2026, 8, 17, 12, 0, 0),
            )

            review_book = Workbook()
            sheet = review_book.active
            sheet.title = LOCAL_REVIEW_SHEET_NAME
            sheet.append(LOCAL_REVIEW_HEADERS)
            sheet.append(
                [
                    subject.test_id,
                    subject.target_name,
                    subject.recipient_email,
                    subject.recipient_type,
                    subject.template_type,
                    subject.language,
                    "Subject",
                    "Body",
                    "Yes",
                    "",
                    "Yes",
                    "",
                    "VALID",
                    "",
                    "2026-08-17T12:05:00+01:00",
                    "",
                ]
            )
            review_book.save(review)
            log.write_text(
                "timestamp,test_id,target_name,recipient_email,recipient_type,template_type,language,subject,status,graph_response,error\n"
                f"2000-01-01T08:55:00+00:00,{subject.test_id},{subject.target_name},{subject.recipient_email},{subject.recipient_type},{subject.template_type},{subject.language},Subject,OUTLOOK_SCHEDULED,Outlook deferred send requested (2000-01-01 09:00),\n",
                encoding="utf-8",
            )

            statuses = load_target_statuses(source, handoff, review, log)
            status = next(item for item in statuses if item.test_id == subject.test_id)

            self.assertEqual("2000-01-01 09:00", status.scheduled_send_at)
            self.assertEqual("OUTLOOK_SCHEDULED", status.last_send_status)
            self.assertEqual("Check scheduled sent", status.next_action)

    def test_status_treats_literal_none_schedule_as_blank(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.xlsx"
            handoff = Path(temp_dir) / "handoff.xlsx"
            review = Path(temp_dir) / "review.xlsx"
            log = Path(temp_dir) / "send_log.csv"

            workbook = load_workbook(config.EXCEL_PATH)
            workbook.save(source)
            subjects, _templates = __import__("outreach.excel_reader", fromlist=["read_workbook"]).read_workbook(source)
            subject = next(item for item in subjects if item.enabled)

            review_book = Workbook()
            sheet = review_book.active
            sheet.title = LOCAL_REVIEW_SHEET_NAME
            sheet.append(LOCAL_REVIEW_HEADERS)
            sheet.append(
                [
                    subject.test_id,
                    subject.target_name,
                    subject.recipient_email,
                    subject.recipient_type,
                    subject.template_type,
                    subject.language,
                    "Subject",
                    "Body",
                    "Yes",
                    "None",
                    "Yes",
                    "",
                    "VALID",
                    "",
                    "2026-08-17T12:05:00+01:00",
                    "",
                ]
            )
            review_book.save(review)
            log.write_text(
                "timestamp,test_id,target_name,recipient_email,recipient_type,template_type,language,subject,scheduled_send_at,status,graph_response,error\n"
                f"2000-01-01T08:55:00+00:00,{subject.test_id},{subject.target_name},{subject.recipient_email},{subject.recipient_type},{subject.template_type},{subject.language},Subject,None,DRAFT_CREATED,,\n",
                encoding="utf-8",
            )

            statuses = load_target_statuses(source, handoff, review, log)
            status = next(item for item in statuses if item.test_id == subject.test_id)

            self.assertEqual("", status.scheduled_send_at)
            self.assertNotEqual("None", status.scheduled_send_at)

    def test_sent_row_stays_done_when_streamlined_review_omits_it(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.xlsx"
            handoff = Path(temp_dir) / "handoff.xlsx"
            review = Path(temp_dir) / "review.xlsx"
            log = Path(temp_dir) / "send_log.csv"

            workbook = load_workbook(config.EXCEL_PATH)
            workbook.save(source)
            subjects, _templates = __import__("outreach.excel_reader", fromlist=["read_workbook"]).read_workbook(source)
            subject = next(item for item in subjects if item.enabled)

            review_book = Workbook()
            sheet = review_book.active
            sheet.title = LOCAL_REVIEW_SHEET_NAME
            sheet.append(LOCAL_REVIEW_HEADERS)
            review_book.save(review)

            email = QueueEmail(
                test_id=subject.test_id,
                target_name=subject.target_name,
                recipient_email=subject.recipient_email,
                recipient_type=subject.recipient_type,
                template_type=subject.template_type,
                language=subject.language,
                subject="Subject",
                body="Body",
                ready_to_send=True,
                validation_status="VALID",
            )
            append_send_log(log, email, "SENT", graph_response="test")

            statuses = load_target_statuses(source, handoff, review, log)
            status = next(item for item in statuses if item.test_id == subject.test_id)

            self.assertEqual("SENT", status.last_send_status)
            self.assertEqual("", status.issues)
            self.assertEqual("Done", status.next_action)

    def test_transient_sending_log_does_not_replace_previous_durable_status(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.xlsx"
            handoff = Path(temp_dir) / "handoff.xlsx"
            review = Path(temp_dir) / "review.xlsx"
            log = Path(temp_dir) / "send_log.csv"

            workbook = load_workbook(config.EXCEL_PATH)
            workbook.save(source)
            subjects, _templates = __import__("outreach.excel_reader", fromlist=["read_workbook"]).read_workbook(source)
            subject = next(item for item in subjects if item.enabled)

            review_book = Workbook()
            sheet = review_book.active
            sheet.title = LOCAL_REVIEW_SHEET_NAME
            sheet.append(LOCAL_REVIEW_HEADERS)
            review_book.save(review)

            email = QueueEmail(
                test_id=subject.test_id,
                target_name=subject.target_name,
                recipient_email=subject.recipient_email,
                recipient_type=subject.recipient_type,
                template_type=subject.template_type,
                language=subject.language,
                subject="Subject",
                body="Body",
                ready_to_send=True,
                validation_status="VALID",
                scheduled_send_at="2000-01-01 09:00",
            )
            append_send_log(log, email, "OUTLOOK_OUTBOX_LOCKED")
            append_send_log(log, email, "SENDING")

            statuses = load_target_statuses(source, handoff, review, log)
            status = next(item for item in statuses if item.test_id == subject.test_id)

            self.assertEqual("OUTLOOK_OUTBOX_LOCKED", status.last_send_status)
            self.assertEqual("Fix outbox send now", status.next_action)

    def test_active_source_has_no_missing_email_rows(self) -> None:
        statuses = load_target_statuses()

        self.assertFalse([item for item in statuses if not item.recipient_email.strip()])

    def test_status_flags_delivery_failed_rows(self) -> None:
        issue = _issue_summary(
            recipient_email="test@example.com",
            enabled="Yes",
            handoff_generated="Yes",
            review_imported="Yes",
            ready_to_send="Yes",
            validation_status="VALID",
            scheduled_send_at="",
            last_send_status="DELIVERY_FAILED",
        )
        next_action = _next_action(
            recipient_email="test@example.com",
            enabled="Yes",
            handoff_generated="Yes",
            review_imported="Yes",
            ready_to_send="Yes",
            validation_status="VALID",
            scheduled_send_at="",
            last_send_status="DELIVERY_FAILED",
        )

        self.assertEqual("Delivery failed", issue)
        self.assertEqual("Fix recipient email", next_action)

    def test_status_flags_refused_to_answer_rows(self) -> None:
        issue = _issue_summary(
            recipient_email="test@example.com",
            enabled="Yes",
            handoff_generated="Yes",
            review_imported="Yes",
            ready_to_send="Yes",
            validation_status="VALID",
            scheduled_send_at="",
            last_send_status="REFUSED_TO_ANSWER",
        )
        next_action = _next_action(
            recipient_email="test@example.com",
            enabled="Yes",
            handoff_generated="Yes",
            review_imported="Yes",
            ready_to_send="Yes",
            validation_status="VALID",
            scheduled_send_at="",
            last_send_status="REFUSED_TO_ANSWER",
        )

        self.assertEqual("Refused to answer", issue)
        self.assertEqual("Keep excluded", next_action)


if __name__ == "__main__":
    unittest.main()

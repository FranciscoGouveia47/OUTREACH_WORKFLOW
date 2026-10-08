from __future__ import annotations

import csv
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

import config
from outreach.excel_reader import read_workbook
from outreach.handoff_workbook import SHEET_NAME, write_handoff_workbook
from outreach.logger import DRAFT_CREATED_STATUS, append_send_log
from outreach.models import HandoffEmail, QueueEmail
from outreach.prompt_packet import build_prompt_packet
from outreach.reset_entry import reset_target_entry
from outreach.review_importer import LOCAL_REVIEW_HEADERS, LOCAL_REVIEW_SHEET_NAME
from outreach.template_engine import render_email


class ResetEntryTests(unittest.TestCase):
    def test_reset_target_entry_removes_local_state_and_deletes_logged_outlook_draft(self) -> None:
        subjects, templates = read_workbook(config.EXCEL_PATH)
        subject = next(item for item in subjects if item.enabled)
        email_message = render_email(subject, templates)

        with tempfile.TemporaryDirectory() as temp_dir:
            handoff = Path(temp_dir) / "handoff.xlsx"
            review = Path(temp_dir) / "review.xlsx"
            log = Path(temp_dir) / "send_log.csv"

            write_handoff_workbook(
                [
                    HandoffEmail(
                        stakeholder=subject,
                        email=email_message,
                        prompt_packet=build_prompt_packet(subject, email_message),
                    )
                ],
                handoff,
                datetime(2026, 8, 17, 12, 0, 0),
            )

            review_book = Workbook()
            sheet = review_book.active
            sheet.title = LOCAL_REVIEW_SHEET_NAME
            sheet.append(LOCAL_REVIEW_HEADERS)
            sheet.append(
                [
                    subject.test_id,
                    subject.target_name,
                    subject.recipient_email,
                    subject.recipient_type,
                    subject.template_type,
                    subject.language,
                    email_message.subject,
                    email_message.body,
                    "Yes",
                    "2026-08-18 11:00",
                    "Yes",
                    "",
                    "VALID",
                    "",
                    "2026-08-17T12:05:00+01:00",
                    "",
                ]
            )
            review_book.save(review)
            review_book.close()

            queue_email = QueueEmail(
                test_id=subject.test_id,
                target_name=subject.target_name,
                recipient_email=subject.recipient_email,
                recipient_type=subject.recipient_type,
                template_type=subject.template_type,
                language=subject.language,
                subject=email_message.subject,
                body=email_message.body,
                ready_to_send=True,
                validation_status="VALID",
            )
            append_send_log(
                log,
                queue_email,
                DRAFT_CREATED_STATUS,
                graph_response="Outlook draft saved (draft-entry-id-123)",
            )

            with (
                patch("config.CHATGPT_HANDOFF_WORKBOOK_PATH", handoff),
                patch("config.LOCAL_REVIEW_WORKBOOK_PATH", review),
                patch("config.SEND_LOG_PATH", log),
                patch("outreach.reset_entry.run_outlook_action") as outlook_action,
            ):
                result = reset_target_entry(subject.test_id)

            self.assertTrue(result.review_removed)
            self.assertTrue(result.handoff_removed)
            self.assertEqual(1, result.log_entries_removed)
            self.assertEqual(1, result.outlook_drafts_deleted)
            outlook_action.assert_called_once()
            self.assertEqual("delete-draft", outlook_action.call_args.args[1])
            self.assertEqual("draft-entry-id-123", outlook_action.call_args.kwargs["draft_entry_id"])

            handoff_book = load_workbook(handoff, data_only=True)
            try:
                ids = [handoff_book[SHEET_NAME].cell(row, 1).value for row in range(2, handoff_book[SHEET_NAME].max_row + 1)]
            finally:
                handoff_book.close()
            self.assertNotIn(subject.test_id, ids)

            review_book = load_workbook(review, data_only=True)
            try:
                self.assertEqual(1, review_book[LOCAL_REVIEW_SHEET_NAME].max_row)
            finally:
                review_book.close()

            with log.open("r", encoding="utf-8", newline="") as handle:
                self.assertEqual([], list(csv.DictReader(handle)))


if __name__ == "__main__":
    unittest.main()

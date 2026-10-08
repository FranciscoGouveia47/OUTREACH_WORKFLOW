from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

import config
from outreach.entry_actions import (
    cancel_outlook_scheduled_for_entry,
    get_base_email,
    get_ai_checked_email,
    regenerate_handoff_for_entries,
    schedule_outlook_for_entry,
    send_outbox_locked_for_entry,
    send_outlook_for_entry,
    toggle_delivery_failed_for_entry,
    toggle_refused_to_answer_for_entry,
)
from outreach.logger import (
    DELIVERY_FAILED_STATUS,
    DELIVERY_FAILURE_CLEARED_STATUS,
    OUTLOOK_OUTBOX_LOCKED_STATUS,
    OUTLOOK_SCHEDULE_CANCELLED_STATUS,
    OUTLOOK_SCHEDULED_STATUS,
    REFUSED_TO_ANSWER_CLEARED_STATUS,
    REFUSED_TO_ANSWER_STATUS,
)
from outreach.models import QueueEmail, Stakeholder
from outreach.review_importer import LOCAL_REVIEW_HEADERS, LOCAL_REVIEW_SHEET_NAME
from outreach.status import TargetStatus


class EntryActionTests(unittest.TestCase):
    def test_get_base_email_for_entry(self) -> None:
        test_id = "COMM-001"
        email = get_base_email(test_id)

        self.assertEqual(test_id, email.stakeholder_id)
        self.assertTrue(email.subject.startswith(config.TEST_SUBJECT_PREFIX))

    def test_get_ai_checked_email_for_entry(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            review_path = Path(temp_dir) / "review.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = LOCAL_REVIEW_SHEET_NAME
            sheet.append(LOCAL_REVIEW_HEADERS)
            sheet.append(
                [
                    "COMM-001",
                    "Target",
                    "test@example.com",
                    "Individual",
                    "General / Fallback",
                    "English",
                    "Subject",
                    "Body",
                    "Yes",
                    "",
                    "Yes",
                    "",
                    "VALID",
                    "",
                    "",
                    "",
                    "None",
                    config.PROMPT_PACKET_VERSION,
                ]
            )
            workbook.save(review_path)
            workbook.close()

            with patch("config.LOCAL_REVIEW_WORKBOOK_PATH", review_path):
                email = get_ai_checked_email("COMM-001")

        self.assertEqual("COMM-001", email.test_id)
        self.assertEqual("VALID", email.validation_status)

    def test_send_entry_requires_ready_flag(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            review_path = Path(temp_dir) / "review.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = LOCAL_REVIEW_SHEET_NAME
            sheet.append(LOCAL_REVIEW_HEADERS)
            sheet.append(
                [
                    "TEST-001",
                    "Target",
                    "test@example.com",
                    "Individual",
                    "General / Fallback",
                    "English",
                    "[DRAFT] Subject",
                    "Body",
                    "No",
                    "",
                    "Yes",
                    "",
                    "VALID",
                    "",
                    "",
                    "",
                ]
            )
            workbook.save(review_path)
            workbook.close()

            with patch("config.LOCAL_REVIEW_WORKBOOK_PATH", review_path):
                with self.assertRaisesRegex(ValueError, "not marked Ready"):
                    send_outlook_for_entry("TEST-001")

    def test_send_now_ignores_future_schedule(self) -> None:
        email = QueueEmail(
            test_id="TEST-001",
            target_name="Target",
            recipient_email="test@example.com",
            recipient_type="Individual",
            template_type="General / Fallback",
            language="English",
            subject="[DRAFT] Subject",
            body="Body",
            ready_to_send=True,
            validation_status="VALID",
            scheduled_send_at="2999-01-01 09:00",
        )

        with (
            patch("outreach.entry_actions.get_ai_checked_email", return_value=email),
            patch("outreach.entry_actions.process_send_queue") as process_send_queue,
        ):
            process_send_queue.return_value = (1, 0, 0)

            result = send_outlook_for_entry("TEST-001")

        self.assertEqual((1, 0, 0), result)
        queued_email = process_send_queue.call_args.args[0][0]
        self.assertEqual("", queued_email.scheduled_send_at)

    def test_schedule_entry_requires_future_schedule(self) -> None:
        email = QueueEmail(
            test_id="TEST-001",
            target_name="Target",
            recipient_email="test@example.com",
            recipient_type="Individual",
            template_type="General / Fallback",
            language="English",
            subject="[DRAFT] Subject",
            body="Body",
            ready_to_send=True,
            validation_status="VALID",
            scheduled_send_at="2999-01-01 09:00",
        )

        with (
            patch("outreach.entry_actions.get_ai_checked_email", return_value=email),
            patch("outreach.entry_actions.process_send_queue") as process_send_queue,
        ):
            process_send_queue.return_value = (1, 0, 0)

            result = schedule_outlook_for_entry("TEST-001")

        self.assertEqual((1, 0, 0), result)
        self.assertEqual("outlook-scheduled", process_send_queue.call_args.kwargs["delivery"])

    def test_cancel_outlook_scheduled_entry_logs_cancel_and_sets_not_ready(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            review_path = Path(temp_dir) / "review.xlsx"
            log_path = Path(temp_dir) / "send_log.csv"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = LOCAL_REVIEW_SHEET_NAME
            sheet.append(LOCAL_REVIEW_HEADERS)
            sheet.append(
                [
                    "TEST-001",
                    "Target",
                    "test@example.com",
                    "Individual",
                    "General / Fallback",
                    "English",
                    "Subject",
                    "Body",
                    "Yes",
                    "2999-01-01 09:00",
                    "Yes",
                    "",
                    "VALID",
                    "",
                    "",
                    "",
                    "None",
                    config.PROMPT_PACKET_VERSION,
                ]
            )
            workbook.save(review_path)
            workbook.close()
            status = TargetStatus(
                test_id="TEST-001",
                target_name="Target",
                recipient_email="test@example.com",
                enabled="Yes",
                handoff_generated="Yes",
                review_imported="Yes",
                ready_to_send="Yes",
                validation_status="VALID",
                scheduled_send_at="2999-01-01 09:00",
                last_send_status=OUTLOOK_SCHEDULED_STATUS,
                last_send_time="",
                template_type="General / Fallback",
                language="English",
            )

            with (
                patch("config.LOCAL_REVIEW_WORKBOOK_PATH", review_path),
                patch("config.SEND_LOG_PATH", log_path),
                patch("outreach.entry_actions.load_target_statuses", return_value=[status]),
                patch("outreach.entry_actions.run_outlook_action", return_value="Outlook scheduled send cancelled (2999-01-01 09:00)") as outlook_action,
            ):
                result = cancel_outlook_scheduled_for_entry("TEST-001")

            self.assertEqual((1, 0, 0), result)
            outlook_action.assert_called_once()
            self.assertIn(OUTLOOK_SCHEDULE_CANCELLED_STATUS, log_path.read_text(encoding="utf-8"))
            updated = __import__("openpyxl", fromlist=["load_workbook"]).load_workbook(review_path, data_only=True)
            try:
                self.assertEqual("No", updated[LOCAL_REVIEW_SHEET_NAME]["I2"].value)
            finally:
                updated.close()

    def test_cancel_outlook_scheduled_entry_skips_when_not_scheduled(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            review_path = Path(temp_dir) / "review.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = LOCAL_REVIEW_SHEET_NAME
            sheet.append(LOCAL_REVIEW_HEADERS)
            sheet.append(
                [
                    "TEST-001",
                    "Target",
                    "test@example.com",
                    "Individual",
                    "General / Fallback",
                    "English",
                    "Subject",
                    "Body",
                    "Yes",
                    "",
                    "Yes",
                    "",
                    "VALID",
                    "",
                    "",
                    "",
                    "None",
                    config.PROMPT_PACKET_VERSION,
                ]
            )
            workbook.save(review_path)
            workbook.close()
            status = TargetStatus(
                test_id="TEST-001",
                target_name="Target",
                recipient_email="test@example.com",
                enabled="Yes",
                handoff_generated="Yes",
                review_imported="Yes",
                ready_to_send="Yes",
                validation_status="VALID",
                scheduled_send_at="",
                last_send_status="",
                last_send_time="",
                template_type="General / Fallback",
                language="English",
            )

            with (
                patch("config.LOCAL_REVIEW_WORKBOOK_PATH", review_path),
                patch("outreach.entry_actions.load_target_statuses", return_value=[status]),
                patch("outreach.entry_actions.run_outlook_action") as outlook_action,
            ):
                result = cancel_outlook_scheduled_for_entry("TEST-001")

            self.assertEqual((0, 1, 0), result)
            outlook_action.assert_not_called()
            updated = __import__("openpyxl", fromlist=["load_workbook"]).load_workbook(review_path, data_only=True)
            try:
                self.assertEqual("No", updated[LOCAL_REVIEW_SHEET_NAME]["I2"].value)
            finally:
                updated.close()

    def test_send_outbox_locked_entry_sends_existing_outbox_item_now(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            log_path = Path(temp_dir) / "send_log.csv"
            email = QueueEmail(
                test_id="TEST-001",
                target_name="Target",
                recipient_email="test@example.com",
                recipient_type="Individual",
                template_type="General / Fallback",
                language="English",
                subject="Subject handed to Outlook",
                body="Body",
                ready_to_send=True,
                validation_status="VALID",
                scheduled_send_at="2000-01-01 09:00",
            )
            from outreach.logger import append_send_log

            append_send_log(log_path, email, OUTLOOK_SCHEDULED_STATUS)
            append_send_log(log_path, email, OUTLOOK_OUTBOX_LOCKED_STATUS)
            status = TargetStatus(
                test_id="TEST-001",
                target_name="Target",
                recipient_email="test@example.com",
                enabled="Yes",
                handoff_generated="Yes",
                review_imported="No",
                ready_to_send="",
                validation_status="",
                scheduled_send_at="2000-01-01 09:00",
                last_send_status=OUTLOOK_OUTBOX_LOCKED_STATUS,
                last_send_time="",
                template_type="General / Fallback",
                language="English",
            )

            with (
                patch("config.SEND_LOG_PATH", log_path),
                patch("outreach.entry_actions.load_target_statuses", return_value=[status]),
                patch(
                    "outreach.entry_actions.run_outlook_action",
                    return_value="Outlook outbox message sent now (2000-01-01 09:00)",
                ) as outlook_action,
            ):
                result = send_outbox_locked_for_entry("TEST-001")

            self.assertEqual((1, 0, 0), result)
            outlook_action.assert_called_once()
            self.assertEqual("send-outbox-now", outlook_action.call_args.args[1])
            self.assertIn("SENT", log_path.read_text(encoding="utf-8"))

    def test_send_outbox_locked_entry_skips_when_not_locked(self) -> None:
        status = TargetStatus(
            test_id="TEST-001",
            target_name="Target",
            recipient_email="test@example.com",
            enabled="Yes",
            handoff_generated="Yes",
            review_imported="Yes",
            ready_to_send="Yes",
            validation_status="VALID",
            scheduled_send_at="",
            last_send_status=OUTLOOK_SCHEDULED_STATUS,
            last_send_time="",
            template_type="General / Fallback",
            language="English",
        )

        with (
            patch("outreach.entry_actions.load_target_statuses", return_value=[status]),
            patch("outreach.entry_actions.run_outlook_action") as outlook_action,
        ):
            result = send_outbox_locked_for_entry("TEST-001")

        self.assertEqual((0, 1, 0), result)
        outlook_action.assert_not_called()

    def test_send_outbox_locked_entry_logs_locked_when_outlook_fails(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            log_path = Path(temp_dir) / "send_log.csv"
            email = QueueEmail(
                test_id="TEST-001",
                target_name="Target",
                recipient_email="test@example.com",
                recipient_type="Individual",
                template_type="General / Fallback",
                language="English",
                subject="Subject handed to Outlook",
                body="Body",
                ready_to_send=True,
                validation_status="VALID",
                scheduled_send_at="2000-01-01 09:00",
            )
            from outreach.logger import append_send_log, latest_status

            append_send_log(log_path, email, OUTLOOK_OUTBOX_LOCKED_STATUS)
            status = TargetStatus(
                test_id="TEST-001",
                target_name="Target",
                recipient_email="test@example.com",
                enabled="Yes",
                handoff_generated="Yes",
                review_imported="No",
                ready_to_send="",
                validation_status="",
                scheduled_send_at="2000-01-01 09:00",
                last_send_status=OUTLOOK_OUTBOX_LOCKED_STATUS,
                last_send_time="",
                template_type="General / Fallback",
                language="English",
            )

            with (
                patch("config.SEND_LOG_PATH", log_path),
                patch("outreach.entry_actions.load_target_statuses", return_value=[status]),
                patch("outreach.entry_actions.run_outlook_action", side_effect=RuntimeError("Outlook refused")),
            ):
                result = send_outbox_locked_for_entry("TEST-001")

            self.assertEqual((0, 0, 1), result)
            self.assertEqual(OUTLOOK_OUTBOX_LOCKED_STATUS, latest_status(log_path, "TEST-001", "test@example.com"))

    def test_regenerate_handoff_skips_delivery_failed_entries(self) -> None:
        subject = Stakeholder(
            source_row=2,
            test_id="TEST-001",
            target_name="Target",
            salutation="Target",
            organisation="Org",
            role="Researcher",
            professional_context="Academic / research",
            sector="Bio-based materials",
            personalisation_fact="Relevant expertise.",
            language="English",
            recipient_email="test@example.com",
            enabled=True,
            recipient_type="Individual",
            template_type="General / Fallback",
            handoff_generated=False,
        )

        with (
            patch("outreach.entry_actions.read_workbook", return_value=([subject], {})),
            patch("outreach.entry_actions.already_sent", return_value=False),
            patch("outreach.entry_actions.has_delivery_failed", return_value=True),
            patch("outreach.entry_actions.render_email") as render_email,
            patch("outreach.entry_actions.upsert_handoff_workbook_rows") as upsert,
        ):
            upsert.return_value = config.CHATGPT_HANDOFF_WORKBOOK_PATH

            generated_count, path = regenerate_handoff_for_entries(["TEST-001"])

        self.assertEqual(0, generated_count)
        self.assertEqual(config.CHATGPT_HANDOFF_WORKBOOK_PATH, path)
        render_email.assert_not_called()
        upsert.assert_called_once()
        self.assertEqual([], upsert.call_args.args[0])

    def test_regenerate_handoff_skips_outlook_scheduled_entries(self) -> None:
        subject = Stakeholder(
            source_row=2,
            test_id="TEST-001",
            target_name="Target",
            salutation="Target",
            organisation="Org",
            role="Researcher",
            professional_context="Academic / research",
            sector="Bio-based materials",
            personalisation_fact="Relevant expertise.",
            language="English",
            recipient_email="test@example.com",
            enabled=True,
            recipient_type="Individual",
            template_type="General / Fallback",
            handoff_generated=False,
        )

        with (
            patch("outreach.entry_actions.read_workbook", return_value=([subject], {})),
            patch("outreach.entry_actions.already_sent", return_value=False),
            patch("outreach.entry_actions.has_outlook_scheduled", return_value=True),
            patch("outreach.entry_actions.has_delivery_failed") as has_delivery_failed,
            patch("outreach.entry_actions.render_email") as render_email,
            patch("outreach.entry_actions.upsert_handoff_workbook_rows") as upsert,
        ):
            upsert.return_value = config.CHATGPT_HANDOFF_WORKBOOK_PATH

            generated_count, path = regenerate_handoff_for_entries(["TEST-001"])

        self.assertEqual(0, generated_count)
        self.assertEqual(config.CHATGPT_HANDOFF_WORKBOOK_PATH, path)
        has_delivery_failed.assert_not_called()
        render_email.assert_not_called()
        upsert.assert_called_once()
        self.assertEqual([], upsert.call_args.args[0])

    def test_regenerate_handoff_skips_refused_to_answer_entries(self) -> None:
        subject = Stakeholder(
            source_row=2,
            test_id="TEST-001",
            target_name="Target",
            salutation="Target",
            organisation="Org",
            role="Researcher",
            professional_context="Academic / research",
            sector="Bio-based materials",
            personalisation_fact="Relevant expertise.",
            language="English",
            recipient_email="test@example.com",
            enabled=True,
            recipient_type="Individual",
            template_type="General / Fallback",
            handoff_generated=False,
        )

        with (
            patch("outreach.entry_actions.read_workbook", return_value=([subject], {})),
            patch("outreach.entry_actions.already_sent", return_value=False),
            patch("outreach.entry_actions.has_delivery_failed", return_value=False),
            patch("outreach.entry_actions.has_refused_to_answer", return_value=True),
            patch("outreach.entry_actions.render_email") as render_email,
            patch("outreach.entry_actions.upsert_handoff_workbook_rows") as upsert,
        ):
            upsert.return_value = config.CHATGPT_HANDOFF_WORKBOOK_PATH

            generated_count, path = regenerate_handoff_for_entries(["TEST-001"])

        self.assertEqual(0, generated_count)
        self.assertEqual(config.CHATGPT_HANDOFF_WORKBOOK_PATH, path)
        render_email.assert_not_called()
        upsert.assert_called_once()
        self.assertEqual([], upsert.call_args.args[0])

    def test_toggle_delivery_failed_logs_manual_flag(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            log_path = Path(temp_dir) / "send_log.csv"
            status = TargetStatus(
                test_id="TEST-001",
                target_name="Target",
                recipient_email="test@example.com",
                enabled="Yes",
                handoff_generated="Yes",
                review_imported="Yes",
                ready_to_send="Yes",
                validation_status="VALID",
                scheduled_send_at="",
                last_send_status="SENT",
                last_send_time="",
                template_type="General / Fallback",
                language="English",
            )
            email = QueueEmail(
                test_id="TEST-001",
                target_name="Target",
                recipient_email="test@example.com",
                recipient_type="Individual",
                template_type="General / Fallback",
                language="English",
                subject="Subject",
                body="Body",
                ready_to_send=True,
                validation_status="VALID",
            )

            with (
                patch("config.SEND_LOG_PATH", log_path),
                patch("outreach.entry_actions.load_target_statuses", return_value=[status]),
                patch("outreach.entry_actions.read_review_email", return_value=email),
            ):
                result = toggle_delivery_failed_for_entry("TEST-001")

            self.assertEqual(DELIVERY_FAILED_STATUS, result)
            self.assertIn(DELIVERY_FAILED_STATUS, log_path.read_text(encoding="utf-8"))

    def test_toggle_delivery_failed_clears_manual_flag(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            log_path = Path(temp_dir) / "send_log.csv"
            status = TargetStatus(
                test_id="TEST-001",
                target_name="Target",
                recipient_email="test@example.com",
                enabled="Yes",
                handoff_generated="Yes",
                review_imported="Yes",
                ready_to_send="Yes",
                validation_status="VALID",
                scheduled_send_at="",
                last_send_status=DELIVERY_FAILED_STATUS,
                last_send_time="",
                template_type="General / Fallback",
                language="English",
            )
            email = QueueEmail(
                test_id="TEST-001",
                target_name="Target",
                recipient_email="test@example.com",
                recipient_type="Individual",
                template_type="General / Fallback",
                language="English",
                subject="Subject",
                body="Body",
                ready_to_send=True,
                validation_status="VALID",
            )

            with (
                patch("config.SEND_LOG_PATH", log_path),
                patch("outreach.entry_actions.load_target_statuses", return_value=[status]),
                patch("outreach.entry_actions.read_review_email", return_value=email),
            ):
                result = toggle_delivery_failed_for_entry("TEST-001")

            self.assertEqual(DELIVERY_FAILURE_CLEARED_STATUS, result)
            self.assertIn(DELIVERY_FAILURE_CLEARED_STATUS, log_path.read_text(encoding="utf-8"))

    def test_toggle_refused_to_answer_logs_manual_flag(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            log_path = Path(temp_dir) / "send_log.csv"
            status = TargetStatus(
                test_id="TEST-001",
                target_name="Target",
                recipient_email="test@example.com",
                enabled="Yes",
                handoff_generated="Yes",
                review_imported="Yes",
                ready_to_send="Yes",
                validation_status="VALID",
                scheduled_send_at="",
                last_send_status="SENT",
                last_send_time="",
                template_type="General / Fallback",
                language="English",
            )
            email = QueueEmail(
                test_id="TEST-001",
                target_name="Target",
                recipient_email="test@example.com",
                recipient_type="Individual",
                template_type="General / Fallback",
                language="English",
                subject="Subject",
                body="Body",
                ready_to_send=True,
                validation_status="VALID",
            )

            with (
                patch("config.SEND_LOG_PATH", log_path),
                patch("outreach.entry_actions.load_target_statuses", return_value=[status]),
                patch("outreach.entry_actions.read_review_email", return_value=email),
            ):
                result = toggle_refused_to_answer_for_entry("TEST-001")

            self.assertEqual(REFUSED_TO_ANSWER_STATUS, result)
            self.assertIn(REFUSED_TO_ANSWER_STATUS, log_path.read_text(encoding="utf-8"))

    def test_toggle_refused_to_answer_clears_manual_flag(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            log_path = Path(temp_dir) / "send_log.csv"
            status = TargetStatus(
                test_id="TEST-001",
                target_name="Target",
                recipient_email="test@example.com",
                enabled="Yes",
                handoff_generated="Yes",
                review_imported="Yes",
                ready_to_send="Yes",
                validation_status="VALID",
                scheduled_send_at="",
                last_send_status=REFUSED_TO_ANSWER_STATUS,
                last_send_time="",
                template_type="General / Fallback",
                language="English",
            )
            email = QueueEmail(
                test_id="TEST-001",
                target_name="Target",
                recipient_email="test@example.com",
                recipient_type="Individual",
                template_type="General / Fallback",
                language="English",
                subject="Subject",
                body="Body",
                ready_to_send=True,
                validation_status="VALID",
            )

            with (
                patch("config.SEND_LOG_PATH", log_path),
                patch("outreach.entry_actions.load_target_statuses", return_value=[status]),
                patch("outreach.entry_actions.read_review_email", return_value=email),
            ):
                result = toggle_refused_to_answer_for_entry("TEST-001")

            self.assertEqual(REFUSED_TO_ANSWER_CLEARED_STATUS, result)
            self.assertIn(REFUSED_TO_ANSWER_CLEARED_STATUS, log_path.read_text(encoding="utf-8"))


if __name__ == "__main__":
    unittest.main()

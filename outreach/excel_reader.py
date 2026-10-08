from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from outreach.models import EmailTemplate, Stakeholder


TEST_SUBJECTS_SHEET = "Test Subjects"
EMAIL_TEMPLATES_SHEET = "Email Templates"


def read_workbook(path: Path) -> tuple[list[Stakeholder], dict[str, EmailTemplate]]:
    if not path.exists():
        raise FileNotFoundError(f"Workbook not found: {path}")

    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        missing_sheets = {
            TEST_SUBJECTS_SHEET,
            EMAIL_TEMPLATES_SHEET,
        }.difference(workbook.sheetnames)
        if missing_sheets:
            raise ValueError(f"Workbook missing sheet(s): {', '.join(sorted(missing_sheets))}")

        return (
            _read_subjects(workbook[TEST_SUBJECTS_SHEET]),
            _read_templates(workbook[EMAIL_TEMPLATES_SHEET]),
        )
    finally:
        workbook.close()


def _read_subjects(sheet: Any) -> list[Stakeholder]:
    rows = _rows_as_dicts(sheet)
    subjects: list[Stakeholder] = []

    for source_row, row in rows:
        if not any(row.values()):
            continue

        subjects.append(
            Stakeholder(
                source_row=source_row,
                test_id=_clean(row.get("Test ID")),
                target_name=_clean(row.get("Target Name")),
                salutation=_normalise_salutation(row.get("Salutation")),
                organisation=_clean(row.get("Organisation")),
                role=_clean(row.get("Role / Position")),
                professional_context=_clean(row.get("Professional Context")),
                sector=_clean(row.get("Relevant Sector / Field")),
                personalisation_fact=_clean(row.get("Personalisation Fact")),
                language=_clean(row.get("Language")),
                recipient_email=_clean(row.get("Test Recipient Email")).lower(),
                enabled=_clean(row.get("Enabled")).lower() == "yes",
                recipient_type=_clean(row.get("Recipient Type")),
                template_type=_clean(row.get("Template Type")),
                handoff_generated=_is_truthy(row.get("ChatGPT Handoff Generated")),
            )
        )

    return subjects


def _read_templates(sheet: Any) -> dict[str, EmailTemplate]:
    rows = _rows_as_dicts(sheet)
    templates: dict[str, EmailTemplate] = {}

    for _, row in rows:
        template_type = _clean(row.get("Template Type"))
        if not template_type:
            continue

        templates[template_type] = EmailTemplate(
            template_type=template_type,
            typical_audience=_clean(row.get("Typical Audience")),
            subject_template=_clean(row.get("Subject Template")),
            body_template=_clean(row.get("Body Template"), preserve_newlines=True),
            usage_notes=_clean(row.get("Usage Notes")),
        )

    return templates


def _rows_as_dicts(sheet: Any) -> list[tuple[int, dict[str, Any]]]:
    headers = [_clean(cell.value) for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    rows: list[tuple[int, dict[str, Any]]] = []

    for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        row = {
            header: value
            for header, value in zip(headers, values)
            if header
        }
        rows.append((row_number, row))

    return rows


def _clean(value: Any, preserve_newlines: bool = False) -> str:
    if value is None:
        return ""
    text = str(value).replace("\xa0", " ")
    if preserve_newlines:
        lines = [" ".join(line.split()) for line in text.splitlines()]
        return "\n".join(lines).strip()
    return " ".join(text.split())


def _normalise_salutation(value: Any) -> str:
    salutation = _clean(value)
    while salutation.lower().startswith("dear "):
        salutation = salutation[5:].strip()
    return salutation.rstrip(",").strip()


def _is_truthy(value: Any) -> bool:
    return _clean(value).lower() in {"yes", "true", "1", "generated"}

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

import config
from outreach.lineage import handoff_id_from_workbook, write_metadata
from outreach.logger import already_sent
from outreach.models import ReviewEmail, Stakeholder
from outreach.validator import unresolved_placeholders
from outreach.workbook_io import save_workbook_atomically


HANDOFF_SHEET_NAME = "ChatGPT Handoff"
LOCAL_REVIEW_SHEET_NAME = "Email Review Queue"
LOCAL_REVIEW_HEADERS = [
    "Test ID",
    "Target Name",
    "Recipient Email",
    "Recipient Type",
    "Template Type",
    "Language",
    "Subject",
    "Email Body",
    "Ready To Send",
    "Scheduled Send At",
    "AI Ready For Human Review",
    "AI Quality Notes",
    "Validation Status",
    "Validation Errors",
    "Imported At",
    "Review Notes",
    "AI Risk Flags",
    "Prompt Version",
]


def import_completed_handoff(
    completed_handoff_path: Path,
    subjects: list[Stakeholder],
    local_review_path: Path,
    imported_at: datetime,
    send_log_path: Path | None = None,
) -> tuple[Path, list[ReviewEmail]]:
    review_emails = read_completed_handoff(completed_handoff_path, subjects)
    if send_log_path is not None:
        review_emails = [
            item
            for item in review_emails
            if not already_sent(
                send_log_path,
                item.stakeholder.test_id,
                item.stakeholder.recipient_email,
            )
        ]
    source_handoff_id = handoff_id_from_workbook(completed_handoff_path)
    write_local_review_queue(review_emails, local_review_path, imported_at, source_handoff_id)
    return local_review_path, review_emails


def read_completed_handoff(
    completed_handoff_path: Path,
    subjects: list[Stakeholder],
) -> list[ReviewEmail]:
    if not completed_handoff_path.exists():
        raise FileNotFoundError(f"Completed handoff workbook not found: {completed_handoff_path}")

    subject_by_id = {subject.test_id: subject for subject in subjects}
    workbook = load_workbook(completed_handoff_path, data_only=True)
    try:
        sheet = (
            workbook[HANDOFF_SHEET_NAME]
            if HANDOFF_SHEET_NAME in workbook.sheetnames
            else workbook.active
        )
        headers = {
            str(sheet.cell(1, column).value).strip(): column
            for column in range(1, sheet.max_column + 1)
            if sheet.cell(1, column).value
        }

        required = [
            "Test ID",
            "Template Type",
            "Language",
            "AI Subject",
            "AI Email Body",
            "AI Quality Notes",
            "AI Ready For Human Review",
        ]
        missing = [header for header in required if header not in headers]
        if missing:
            raise ValueError(f"Completed handoff workbook missing column(s): {', '.join(missing)}")

        review_emails: list[ReviewEmail] = []
        for row in range(2, sheet.max_row + 1):
            test_id = _cell(sheet, row, headers["Test ID"])
            if not test_id:
                continue

            stakeholder = subject_by_id.get(test_id)
            if stakeholder is None:
                review_emails.append(
                    ReviewEmail(
                        stakeholder=_unknown_stakeholder(test_id),
                        subject=_cell(sheet, row, headers["AI Subject"]),
                        body=_cell(sheet, row, headers["AI Email Body"]),
                        ai_ready_for_human_review=_cell(sheet, row, headers["AI Ready For Human Review"]),
                        ai_quality_notes=_cell(sheet, row, headers["AI Quality Notes"]),
                        ai_risk_flags=_optional_cell(sheet, row, headers, "AI Risk Flags") or "Unknown Test ID",
                        prompt_version=_optional_cell(sheet, row, headers, "Prompt Version") or config.PROMPT_PACKET_VERSION,
                        validation_status="INVALID",
                        validation_errors="Unknown Test ID in completed handoff workbook.",
                    )
                )
                continue

            subject = _cell(sheet, row, headers["AI Subject"])
            body = _cell(sheet, row, headers["AI Email Body"])
            ai_ready = _cell(sheet, row, headers["AI Ready For Human Review"])
            notes = _cell(sheet, row, headers["AI Quality Notes"])
            risk_flags = _optional_cell(sheet, row, headers, "AI Risk Flags")
            prompt_version = _optional_cell(sheet, row, headers, "Prompt Version") or config.PROMPT_PACKET_VERSION
            validation_errors = _validate_review_email(
                stakeholder,
                subject,
                body,
                ai_ready,
                prompt_version,
            )
            validation_status = "VALID" if not validation_errors else "INVALID"

            review_emails.append(
                ReviewEmail(
                    stakeholder=stakeholder,
                    subject=subject,
                    body=body,
                    ai_ready_for_human_review=ai_ready,
                    ai_quality_notes=notes,
                    ai_risk_flags=_merge_risk_flags(risk_flags, _derive_risk_flags(stakeholder, subject, body, notes)),
                    prompt_version=prompt_version,
                    validation_status=validation_status,
                    validation_errors="; ".join(validation_errors),
                )
            )

        return review_emails
    finally:
        workbook.close()


def write_local_review_queue(
    review_emails: list[ReviewEmail],
    local_review_path: Path,
    imported_at: datetime,
    source_handoff_id: str = "",
) -> Path:
    local_review_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = LOCAL_REVIEW_SHEET_NAME
    sheet.append(LOCAL_REVIEW_HEADERS)

    for item in review_emails:
        stakeholder = item.stakeholder
        sheet.append(
            [
                stakeholder.test_id,
                stakeholder.target_name,
                stakeholder.recipient_email,
                stakeholder.recipient_type,
                stakeholder.template_type,
                stakeholder.language,
                item.subject,
                item.body,
                "No",
                "",
                item.ai_ready_for_human_review,
                item.ai_quality_notes,
                item.validation_status,
                item.validation_errors,
                imported_at.isoformat(timespec="seconds"),
                "",
                item.ai_risk_flags,
                item.prompt_version,
            ]
        )

    _format_sheet(sheet)
    write_metadata(
        workbook,
        {
            "file_type": "local_review_queue",
            "source_handoff_id": source_handoff_id,
            "imported_at": imported_at.isoformat(timespec="seconds"),
        },
    )
    save_workbook_atomically(workbook, local_review_path)
    workbook.close()
    return local_review_path


def _validate_review_email(
    stakeholder: Stakeholder,
    subject: str,
    body: str,
    ai_ready_for_human_review: str = "",
    prompt_version: str = "",
) -> list[str]:
    errors: list[str] = []
    if prompt_version and prompt_version != config.PROMPT_PACKET_VERSION:
        errors.append(
            f"Prompt Version is stale: {prompt_version} "
            f"(current: {config.PROMPT_PACKET_VERSION})"
        )
    if not subject:
        errors.append("AI Subject is blank")
    if not body:
        errors.append("AI Email Body is blank")
    if ai_ready_for_human_review and ai_ready_for_human_review.strip().lower() not in {"yes", "y", "true", "1"}:
        errors.append("AI Ready For Human Review is not Yes")
    if subject and not subject.startswith(config.TEST_SUBJECT_PREFIX):
        errors.append("AI Subject is missing test prefix")
    if stakeholder.recipient_email and stakeholder.recipient_email.lower() in f"{subject}\n{body}".lower():
        errors.append("AI output includes the real recipient email address")
    for placeholder in unresolved_placeholders(f"{subject}\n{body}"):
        errors.append(f"Unresolved placeholder: {placeholder}")
    for value in config.LOCKED_VALUES:
        if value and value not in f"{subject}\n{body}":
            errors.append(f"Missing locked value: {value}")
    if body and len(body) < 80:
        errors.append("AI Email Body is unusually short")
    if body and len(body) > 4000:
        errors.append("AI Email Body is unusually long")
    if stakeholder.recipient_type == "Organisation" and not _has_organisation_routing_language(body):
        errors.append("Organisation email may be missing routing language")
    errors.extend(_recipient_facing_source_artifact_errors(body))
    errors.extend(_localisation_errors(stakeholder, body))
    return errors


def _derive_risk_flags(
    stakeholder: Stakeholder,
    subject: str,
    body: str,
    notes: str,
) -> list[str]:
    flags: list[str] = []
    combined = f"{subject}\n{body}\n{notes}".lower()
    if stakeholder.recipient_type == "Organisation" and not _has_organisation_routing_language(body):
        flags.append("organisation routing")
    if stakeholder.language == "Portuguese" and body and not _looks_portuguese(body):
        flags.append("translation uncertainty")
    if any(word in combined for word in ("invent", "unsupported", "uncertain", "weak", "unsafe")):
        flags.append("manual review")
    return flags


def _localisation_errors(stakeholder: Stakeholder, body: str) -> list[str]:
    if not body:
        return []

    errors: list[str] = []
    lowered = body.strip().lower()
    casual_greetings = {
        "Portuguese": ("olá", "ola"),
        "Portuguese (Brazil)": ("olá", "ola"),
        "Spanish": ("hola",),
        "Italian": ("ciao",),
        "French": ("salut",),
        "German": ("hallo",),
    }
    for greeting in casual_greetings.get(stakeholder.language, ()):
        if lowered.startswith(greeting):
            errors.append("Greeting appears too casual for professional outreach")
            break

    if stakeholder.recipient_type == "Organisation":
        singular_organisation_phrases = {
            "Portuguese": ("se preferir ", "se não quiser ", "caso prefira "),
            "Portuguese (Brazil)": ("se preferir ", "se não quiser ", "caso prefira "),
            "Spanish": ("si prefiere ", "si no desea ", "en caso de que prefiera "),
            "Italian": ("se preferisce ", "se non desidera ", "qualora preferisca "),
        }
        for phrase in singular_organisation_phrases.get(stakeholder.language, ()):
            if phrase in lowered:
                errors.append("Organisation recipient uses singular-person wording")
                break

    return errors


def _recipient_facing_source_artifact_errors(body: str) -> list[str]:
    if not body:
        return []

    lowered = body.lower()
    source_artifact_markers = [
        "the row provides",
        "the workbook",
        "the source describes",
        "the source lists",
        "the university directly lists",
        "directly lists his expertise",
        "directly lists her expertise",
        "directly lists their expertise",
        "for that reason, your perspective would be valuable",
        "aligns closely with the areas explored in the consultation",
    ]
    if any(marker in lowered for marker in source_artifact_markers):
        return ["AI output describes source evidence or contact-selection rationale"]
    return []


def _merge_risk_flags(ai_flags: str, derived_flags: list[str]) -> str:
    values: list[str] = []
    for raw_value in [ai_flags, ", ".join(derived_flags)]:
        for part in raw_value.replace(";", ",").split(","):
            value = part.strip()
            if not value or value.lower() == "none":
                continue
            if value.lower() not in {item.lower() for item in values}:
                values.append(value)
    return ", ".join(values) if values else "None"


def _looks_portuguese(text: str) -> bool:
    lowered = f" {text.lower()} "
    markers = [
        " de ",
        " que ",
        " para ",
        " projeto",
        " inquerito",
        " inquérito",
        " questionario",
        " questionário",
        " participação",
        " cumprimentos",
        " obrigado",
        " obrigada",
    ]
    return sum(1 for marker in markers if marker in lowered) >= 3


def _has_organisation_routing_language(text: str) -> bool:
    lowered = text.lower()
    markers = [
        "colleague",
        "team",
        "shared with",
        "share this invitation",
        "best placed",
        "colega",
        "equipa",
        "partilhar",
        "encaminhar",
    ]
    return any(marker in lowered for marker in markers)


def _cell(sheet, row: int, column: int) -> str:
    value = sheet.cell(row, column).value
    if value is None:
        return ""
    return str(value).strip()


def _optional_cell(sheet, row: int, headers: dict[str, int], header: str) -> str:
    if header not in headers:
        return ""
    return _cell(sheet, row, headers[header])


def _unknown_stakeholder(test_id: str) -> Stakeholder:
    return Stakeholder(
        source_row=0,
        test_id=test_id,
        target_name="",
        salutation="",
        organisation="",
        role="",
        professional_context="",
        sector="",
        personalisation_fact="",
        language="",
        recipient_email="",
        enabled=False,
        recipient_type="",
        template_type="",
        handoff_generated=False,
    )


def _format_sheet(sheet) -> None:
    sheet.freeze_panes = "A2"
    widths = {
        "A": 12,
        "B": 24,
        "C": 32,
        "D": 18,
        "E": 24,
        "F": 14,
        "G": 60,
        "H": 90,
        "I": 16,
        "J": 26,
        "K": 45,
        "L": 18,
        "M": 50,
        "N": 22,
        "O": 35,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for cell in sheet[1]:
        cell.font = Font(bold=True)

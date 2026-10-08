from __future__ import annotations

import shutil
import tempfile
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

import config
from outreach.excel_reader import TEST_SUBJECTS_SHEET
from outreach.models import Stakeholder


def mark_handoffs_generated(
    workbook_path: Path,
    handoff_subjects: list[Stakeholder],
    generated_at: datetime,
    handoff_workbook_path: Path,
) -> Path | None:
    if not handoff_subjects:
        return None

    backup_path = _backup_workbook(workbook_path, generated_at)
    workbook = load_workbook(workbook_path)
    sheet = workbook[TEST_SUBJECTS_SHEET]

    generated_column = _ensure_column(sheet, config.HANDOFF_GENERATED_COLUMN)
    generated_at_column = _ensure_column(sheet, config.HANDOFF_GENERATED_AT_COLUMN)
    handoff_column = _ensure_column(sheet, config.HANDOFF_WORKBOOK_COLUMN)

    for subject in handoff_subjects:
        sheet.cell(subject.source_row, generated_column, "Yes")
        sheet.cell(subject.source_row, generated_at_column, generated_at.isoformat(timespec="seconds"))
        sheet.cell(subject.source_row, handoff_column, str(handoff_workbook_path))

    workbook.save(workbook_path)
    workbook.close()
    return backup_path


def mark_reviews_imported(
    workbook_path: Path,
    imported_subjects: list[Stakeholder],
    imported_at: datetime,
) -> Path | None:
    if not imported_subjects:
        return None

    backup_path = _backup_workbook(workbook_path, imported_at)
    workbook = load_workbook(workbook_path)
    sheet = workbook[TEST_SUBJECTS_SHEET]

    imported_column = _ensure_column(sheet, config.REVIEW_IMPORTED_COLUMN)
    imported_at_column = _ensure_column(sheet, config.REVIEW_IMPORTED_AT_COLUMN)

    for subject in imported_subjects:
        sheet.cell(subject.source_row, imported_column, "Yes")
        sheet.cell(subject.source_row, imported_at_column, imported_at.isoformat(timespec="seconds"))

    workbook.save(workbook_path)
    workbook.close()
    return backup_path


def assert_workbooks_writable(*paths: Path) -> None:
    for path in paths:
        if path.exists():
            try:
                with path.open("r+b"):
                    pass
            except PermissionError as exc:
                raise PermissionError(
                    f"Workbook is not writable, possibly open in Excel: {path}"
                ) from exc
            continue

        path.parent.mkdir(parents=True, exist_ok=True)
        try:
            with tempfile.TemporaryFile(dir=path.parent):
                pass
        except PermissionError as exc:
            raise PermissionError(f"Directory is not writable: {path.parent}") from exc


def _backup_workbook(workbook_path: Path, generated_at: datetime) -> Path:
    backup_dir = workbook_path.parent / "backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    backup_path = backup_dir / (
        f"{workbook_path.stem}.backup-{generated_at.strftime('%Y%m%d-%H%M%S')}{workbook_path.suffix}"
    )
    shutil.copy2(workbook_path, backup_path)
    return backup_path


def _ensure_column(sheet, header: str) -> int:
    for cell in sheet[1]:
        if cell.value == header:
            return cell.column

    column = sheet.max_column + 1
    sheet.cell(1, column, header)
    return column

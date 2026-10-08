from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


@dataclass(frozen=True)
class WorkbookMigrationResult:
    path: Path
    changed: bool
    added_columns: tuple[str, ...]


def ensure_sheet_columns(
    workbook_path: Path,
    sheet_name: str,
    required_headers: list[str],
) -> WorkbookMigrationResult:
    if not workbook_path.exists():
        return WorkbookMigrationResult(workbook_path, changed=False, added_columns=())

    workbook = load_workbook(workbook_path)
    try:
        if sheet_name not in workbook.sheetnames:
            return WorkbookMigrationResult(workbook_path, changed=False, added_columns=())

        sheet = workbook[sheet_name]
        headers = {
            str(sheet.cell(1, column).value).strip(): column
            for column in range(1, sheet.max_column + 1)
            if sheet.cell(1, column).value
        }
        added: list[str] = []
        for header in required_headers:
            if header in headers:
                continue
            sheet.cell(1, sheet.max_column + 1, header)
            added.append(header)

        if added:
            workbook.save(workbook_path)
        return WorkbookMigrationResult(workbook_path, changed=bool(added), added_columns=tuple(added))
    finally:
        workbook.close()

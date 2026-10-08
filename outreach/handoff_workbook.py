from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font

import config
from outreach.handoff_rules import CONTACT_FIT_HANDOFF_RULES
from outreach.lineage import handoff_id_from_sheet, write_metadata
from outreach.models import HandoffEmail


SHEET_NAME = "ChatGPT Handoff"
INSTRUCTIONS_SHEET_NAME = "Instructions"
HEADERS = [
    "Test ID",
    "Target Name",
    "Organisation",
    "Role / Position",
    "Professional Context",
    "Relevant Sector / Field",
    "Personalisation Fact",
    "Language",
    "Recipient Type",
    "Template Type",
    "Base Subject",
    "Base Email Body",
    "ChatGPT Prompt Packet",
    "AI Subject",
    "AI Email Body",
    "AI Quality Notes",
    "AI Ready For Human Review",
    "Generated At",
    "AI Risk Flags",
    "Prompt Version",
]

GENERAL_INSTRUCTIONS = f"""You are assisting with preparation of stakeholder outreach emails for this campaign: {config.CAMPAIGN_NAME}.

This workbook contains one row per target in the "ChatGPT Handoff" sheet.

Your task:
1. For each row, read the "ChatGPT Prompt Packet" cell.
2. Use the base subject and base email as the starting point.
3. Customize the email using only the stakeholder context provided in that row.
4. Fill only these columns:
   - AI Subject
   - AI Email Body
   - AI Quality Notes
   - AI Ready For Human Review
   - AI Risk Flags

Rules:
- Do not invent stakeholder facts, credentials, publications, or expertise.
- Treat the workbook row as the only source available in this handoff; do not
  add new facts unless externally verified.
- If you add any stakeholder-specific fact, role, publication, credential, or
  expertise beyond what is already in the workbook row, verify it externally
  first. If external checking is not available, do not add new facts; instead
  rewrite using only the supplied row context and flag any uncertainty in AI
  Quality Notes or AI Risk Flags.
- Do not use or request real recipient email addresses.
- Preserve the campaign URLs, company URLs, reference codes, sender name, sender position, sender organisation, and any subject prefix already present.
- Keep the email concise, professional, and suitable for human review.
- Vary the email structure and phrasing naturally across rows so messages do not
  read like mechanically repeated bulk mail. Use only truthful, stakeholder-
  relevant variation, such as changing sentence order, combining or splitting
  short paragraphs, and choosing different natural transitions. Do not use
  deceptive tactics, hidden text, obfuscation, misleading subjects, or any
  wording whose only purpose is to evade mail filtering.
- Use the requested language. If the requested language is Portuguese, use European Portuguese; if it is Portuguese (Brazil), use Brazilian Portuguese.
- If a fact is weak, irrelevant, humorous, or unsafe to use, omit it and mention that in AI Quality Notes.
- If the row is for an organisation/shared inbox, include careful routing language and do not imply the inbox owner is the respondent.
- Do not mechanically paste the Organisation field into phrases such as "your work at [Organisation]". Many organisation values include departments, product lines, regions, slashes, commas, or combined affiliations that sound unnatural if inserted verbatim. Rewrite affiliation details into natural professional prose before using them. Examples:
  - Bad: "Your work at Acme Group, Procurement / EMEA Supplier Operations..."
    Better: "Given your work on supplier operations across EMEA..."
  - Bad: "Your work at Northstar Operations Ltd, Manufacturing / Supplier onboarding..."
    Better: "Given your responsibility for supplier onboarding in a manufacturing context..."
  - Bad: "Your work at BrightPath Software / Partner ecosystem / Integrations..."
    Better: "Given your team's focus on partner integrations..."
  When unsure how to phrase a complex affiliation, use the role, field, or personalisation fact instead of forcing the raw organisation string into the sentence.
- If the base email has mistranslations, awkward wording, inconsistencies, or unsupported claims, correct them.
- Pay close attention to capital letters that appear unnecessarily in the middle
  of sentences. Use normal sentence case unless the capitalisation is required
  for a proper noun, official project name, acronym, title, grant label, URL, or
  other locked value.
- Set "AI Ready For Human Review" to "Yes" only when the result is coherent and complete enough for a human reviewer to inspect.
- Set "AI Risk Flags" to comma-separated short labels when the email has risks such as weak personalisation, invented-fact risk, translation uncertainty, sensitive wording, or organisation-routing uncertainty. Use "None" when no risk stands out.
- Return an updated Excel workbook with the same sheets and columns.

{CONTACT_FIT_HANDOFF_RULES}

Global localisation requirements:
1. Formality preservation: maintain the same professional and formal register as the base template in every target language.
2. Greeting localisation: preserve the function and formality of the source greeting. Use a natural professional salutation appropriate to the recipient type, target language and country. Do not downgrade a formal greeting to a casual one such as "Hello", "Hi", "Olá", or equivalents unless the source template itself is informal.
3. Natural localisation: treat the source template as a definition of meaning, facts, structure and tone, not sentence-level syntax. Rewrite naturally in the target language rather than translating sentence by sentence.
4. Localise, do not phrase-map: examples are guidance only. Do not apply fixed phrase replacements mechanically; choose wording a native professional speaker would normally use.
5. Recipient classification: before drafting, classify the target as an individual recipient, organisation/team/shared inbox, or uncertain. Use only the available stakeholder data.
6. Recipient consistency is mandatory: maintain the selected recipient perspective throughout the greeting, pronouns, verb forms, possessives, requests, routing sentence, opt-out wording and closing. Do not address a team collectively in the greeting and later switch to singular-person wording.
7. Organisation/team recipients: use plural, collective or naturally impersonal language throughout. If recipient type is genuinely uncertain, use neutral or impersonal constructions that avoid assuming one specific individual wherever the target language permits.
8. Regional variant: use the requested regional language variant and its professional conventions. For Portuguese, use European Portuguese.
9. Native-language QA: before marking an email ready for human review, reread the whole email as if it had originally been written by a native professional speaker in the target language. If any sentence sounds translated, unnatural, overly literal, overly bureaucratic or inconsistent in recipient perspective, rewrite it before completion.
"""


def write_handoff_workbook(
    handoff_emails: list[HandoffEmail],
    handoff_path: Path,
    generated_at: datetime,
) -> Path:
    handoff_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    instructions_sheet = workbook.active
    instructions_sheet.title = INSTRUCTIONS_SHEET_NAME
    _write_instructions(instructions_sheet)

    sheet = workbook.create_sheet(SHEET_NAME)
    sheet.title = SHEET_NAME
    sheet.append(HEADERS)

    for item in handoff_emails:
        sheet.append(_handoff_row_values(item, generated_at))

    _format_sheet(sheet)
    handoff_id = handoff_id_from_sheet(sheet)
    write_metadata(
        workbook,
        {
            "file_type": "chatgpt_handoff",
            "handoff_id": handoff_id,
            "generated_at": generated_at.isoformat(timespec="seconds"),
        },
    )
    workbook.save(handoff_path)
    workbook.close()
    return handoff_path


def upsert_handoff_workbook_rows(
    handoff_emails: list[HandoffEmail],
    handoff_path: Path,
    generated_at: datetime,
) -> Path:
    if not handoff_path.exists():
        return write_handoff_workbook(handoff_emails, handoff_path, generated_at)

    workbook = load_workbook(handoff_path)
    try:
        if INSTRUCTIONS_SHEET_NAME not in workbook.sheetnames:
            instructions_sheet = workbook.create_sheet(INSTRUCTIONS_SHEET_NAME, 0)
            _write_instructions(instructions_sheet)
        if SHEET_NAME not in workbook.sheetnames:
            sheet = workbook.create_sheet(SHEET_NAME)
            sheet.append(HEADERS)
        else:
            sheet = workbook[SHEET_NAME]
            headers = _headers(sheet)
            for header in HEADERS:
                if header not in headers:
                    sheet.cell(1, sheet.max_column + 1, header)
                    headers[header] = sheet.max_column

        replace_ids = {item.stakeholder.test_id for item in handoff_emails}
        headers = _headers(sheet)
        if "Test ID" in headers:
            for row in range(sheet.max_row, 1, -1):
                if str(sheet.cell(row, headers["Test ID"]).value or "").strip() in replace_ids:
                    sheet.delete_rows(row, 1)

        for item in handoff_emails:
            sheet.append(_handoff_row_values(item, generated_at))

        _format_sheet(sheet)
        handoff_id = handoff_id_from_sheet(sheet)
        write_metadata(
            workbook,
            {
                "file_type": "chatgpt_handoff",
                "handoff_id": handoff_id,
                "generated_at": generated_at.isoformat(timespec="seconds"),
            },
        )
        workbook.save(handoff_path)
    finally:
        workbook.close()
    return handoff_path


def read_prompt_packet(handoff_path: Path, test_id: str) -> str:
    if not handoff_path.exists():
        raise FileNotFoundError(f"Handoff workbook not found: {handoff_path}")

    workbook = load_workbook(handoff_path, data_only=True)
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise ValueError(f"Handoff workbook missing sheet: {SHEET_NAME}")
        sheet = workbook[SHEET_NAME]
        headers = _headers(sheet)
        for header in ("Test ID", "ChatGPT Prompt Packet"):
            if header not in headers:
                raise ValueError(f"Handoff workbook missing column: {header}")
        for row in range(2, sheet.max_row + 1):
            if str(sheet.cell(row, headers["Test ID"]).value or "").strip() == test_id:
                return str(sheet.cell(row, headers["ChatGPT Prompt Packet"]).value or "").strip()
        raise ValueError(f"Test ID not found in handoff workbook: {test_id}")
    finally:
        workbook.close()


def read_handoff_test_ids(handoff_path: Path) -> set[str]:
    if not handoff_path.exists():
        return set()

    workbook = load_workbook(handoff_path, data_only=True)
    try:
        if SHEET_NAME not in workbook.sheetnames:
            return set()
        sheet = workbook[SHEET_NAME]
        headers = _headers(sheet)
        if "Test ID" not in headers:
            return set()
        return {
            str(sheet.cell(row, headers["Test ID"]).value or "").strip()
            for row in range(2, sheet.max_row + 1)
            if str(sheet.cell(row, headers["Test ID"]).value or "").strip()
        }
    finally:
        workbook.close()


def remove_handoff_email(handoff_path: Path, test_id: str) -> bool:
    if not handoff_path.exists():
        return False

    workbook = load_workbook(handoff_path)
    try:
        if SHEET_NAME not in workbook.sheetnames:
            return False
        sheet = workbook[SHEET_NAME]
        headers = _headers(sheet)
        if "Test ID" not in headers:
            return False

        removed = False
        for row in range(sheet.max_row, 1, -1):
            if str(sheet.cell(row, headers["Test ID"]).value or "").strip() == test_id:
                sheet.delete_rows(row, 1)
                removed = True

        if removed:
            write_metadata(
                workbook,
                {
                    "file_type": "chatgpt_handoff",
                    "handoff_id": handoff_id_from_sheet(sheet),
                },
            )
            workbook.save(handoff_path)
        return removed
    finally:
        workbook.close()


def _handoff_row_values(item: HandoffEmail, generated_at: datetime) -> list[str]:
    stakeholder = item.stakeholder
    email = item.email
    return [
        stakeholder.test_id,
        stakeholder.target_name,
        stakeholder.organisation,
        stakeholder.role,
        stakeholder.professional_context,
        stakeholder.sector,
        stakeholder.personalisation_fact,
        stakeholder.language,
        stakeholder.recipient_type,
        email.template_type,
        email.subject,
        email.body,
        item.prompt_packet,
        "",
        "",
        "",
        "",
        generated_at.isoformat(timespec="seconds"),
        "",
        config.PROMPT_PACKET_VERSION,
    ]


def _headers(sheet) -> dict[str, int]:
    return {
        str(sheet.cell(1, column).value).strip(): column
        for column in range(1, sheet.max_column + 1)
        if sheet.cell(1, column).value
    }


def _write_instructions(sheet) -> None:
    sheet["A1"] = "General ChatGPT Instructions"
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A2"] = GENERAL_INSTRUCTIONS
    sheet["A2"].alignment = Alignment(vertical="top", wrap_text=True)
    sheet.column_dimensions["A"].width = 120
    sheet.row_dimensions[2].height = 360


def _format_sheet(sheet) -> None:
    sheet.freeze_panes = "A2"
    widths = {
        "A": 12,
        "B": 24,
        "C": 24,
        "D": 24,
        "E": 24,
        "F": 26,
        "G": 44,
        "H": 14,
        "I": 18,
        "J": 24,
        "K": 60,
        "L": 85,
        "M": 100,
        "N": 60,
        "O": 85,
        "P": 45,
        "Q": 24,
        "R": 22,
        "S": 28,
        "T": 18,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for cell in sheet[1]:
        cell.font = Font(bold=True)

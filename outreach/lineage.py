from __future__ import annotations

import hashlib
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook


METADATA_SHEET_NAME = "_OUTREACH_METADATA"
HANDOFF_SHEET_NAME = "ChatGPT Handoff"
HANDOFF_ID_COLUMNS = [
    "Test ID",
    "Target Name",
    "Organisation",
    "Role / Position",
    "Professional Context",
    "Relevant Sector / Field",
    "Personalisation Fact",
    "Language",
    "Recipient Type",
    "Template Type",
    "Base Subject",
    "Base Email Body",
    "ChatGPT Prompt Packet",
]


def handoff_id_from_sheet(sheet) -> str:
    row_values = list(sheet.iter_rows(values_only=True))
    if not row_values:
        raise ValueError("Handoff workbook sheet is empty.")

    headers = {
        _value(header).strip(): index
        for index, header in enumerate(row_values[0])
        if _value(header).strip()
    }
    missing = [header for header in HANDOFF_ID_COLUMNS if header not in headers]
    if missing:
        raise ValueError(f"Handoff workbook missing fingerprint column(s): {', '.join(missing)}")

    rows: list[dict[str, str]] = []
    for row in row_values[1:]:
        test_id = _row_cell(row, headers["Test ID"])
        if not test_id:
            continue
        rows.append(
            {
                header: _row_cell(row, headers[header])
                for header in HANDOFF_ID_COLUMNS
            }
        )

    payload = json.dumps(rows, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def handoff_id_from_workbook(path: Path) -> str:
    if not path.exists():
        raise FileNotFoundError(f"Handoff workbook not found: {path}")
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        if HANDOFF_SHEET_NAME not in workbook.sheetnames:
            raise ValueError(f"Workbook missing sheet: {HANDOFF_SHEET_NAME}")
        return handoff_id_from_sheet(workbook[HANDOFF_SHEET_NAME])
    finally:
        workbook.close()


def write_metadata(workbook: Workbook, values: dict[str, str]) -> None:
    if METADATA_SHEET_NAME in workbook.sheetnames:
        sheet = workbook[METADATA_SHEET_NAME]
        sheet.delete_rows(1, sheet.max_row)
    else:
        sheet = workbook.create_sheet(METADATA_SHEET_NAME)

    sheet.append(["Key", "Value"])
    for key, value in values.items():
        sheet.append([key, value])
    sheet.sheet_state = "hidden"


def read_metadata(path: Path) -> dict[str, str]:
    if not path.exists():
        return {}
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        if METADATA_SHEET_NAME not in workbook.sheetnames:
            return {}
        sheet = workbook[METADATA_SHEET_NAME]
        values: dict[str, str] = {}
        for row in sheet.iter_rows(min_row=2, values_only=True):
            key = _row_cell(row, 0)
            if key:
                values[key] = _row_cell(row, 1)
        return values
    finally:
        workbook.close()


def lineage_status(handoff_path: Path, review_path: Path) -> tuple[str, str]:
    if not handoff_path.exists():
        return "NO_HANDOFF", "No current ChatGPT handoff workbook found."
    if not review_path.exists():
        return "NO_REVIEW", "No local review queue found."

    try:
        current_handoff_id = handoff_id_from_workbook(handoff_path)
    except Exception as exc:
        return "HANDOFF_INVALID", str(exc)

    review_metadata = read_metadata(review_path)
    review_handoff_id = review_metadata.get("source_handoff_id", "")
    if not review_handoff_id:
        return "UNKNOWN", "Review queue has no lineage metadata; re-import a completed handoff."

    if review_handoff_id != current_handoff_id:
        return "MISMATCH", "Review queue was not imported from the current handoff workbook."

    return "OK", "Review queue matches the current handoff workbook."


def _cell(sheet, row: int, column: int) -> str:
    value = sheet.cell(row, column).value
    return _value(value)


def _row_cell(row: tuple, index: int) -> str:
    if index >= len(row):
        return ""
    return _value(row[index])


def _value(value) -> str:
    if value is None:
        return ""
    return str(value).strip()

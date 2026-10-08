from __future__ import annotations

import csv
import sys
from dataclasses import dataclass
from pathlib import Path

import config
from outreach.excel_reader import read_workbook
from outreach.handoff_workbook import read_handoff_test_ids
from outreach.lineage import lineage_status
from outreach.logger import TRANSIENT_STATUSES, clean_scheduled_send_at, scheduled_send_at_from_response
from outreach.review_importer import LOCAL_REVIEW_SHEET_NAME

csv.field_size_limit(sys.maxsize)


@dataclass(frozen=True)
class TargetStatus:
    test_id: str
    target_name: str
    recipient_email: str
    enabled: str
    handoff_generated: str
    review_imported: str
    ready_to_send: str
    validation_status: str
    scheduled_send_at: str
    last_send_status: str
    last_send_time: str
    template_type: str
    language: str
    ai_risk_flags: str = ""
    prompt_version: str = ""
    issues: str = ""
    next_action: str = ""


@dataclass(frozen=True)
class FileLineageStatus:
    status: str
    message: str


def load_target_statuses(
    source_workbook_path: Path = config.EXCEL_PATH,
    handoff_workbook_path: Path = config.CHATGPT_HANDOFF_WORKBOOK_PATH,
    review_workbook_path: Path = config.LOCAL_REVIEW_WORKBOOK_PATH,
    send_log_path: Path = config.SEND_LOG_PATH,
) -> list[TargetStatus]:
    subjects, _ = read_workbook(source_workbook_path)
    handoff_ids = read_handoff_test_ids(handoff_workbook_path)
    review_rows = _read_review_rows(review_workbook_path)
    latest_log = _read_latest_log_status(send_log_path)

    statuses: list[TargetStatus] = []
    for subject in subjects:
        review = review_rows.get(subject.test_id, {})
        recipient = review.get("Recipient Email") or subject.recipient_email
        log = latest_log.get((subject.test_id, recipient.lower()), {})
        scheduled_send_at = (
            clean_scheduled_send_at(review.get(config.SCHEDULED_SEND_AT_COLUMN, ""))
            or clean_scheduled_send_at(log.get("scheduled_send_at", ""))
        )

        statuses.append(
            TargetStatus(
                test_id=subject.test_id,
                target_name=subject.target_name,
                recipient_email=recipient,
                enabled="Yes" if subject.enabled else "No",
                handoff_generated="Yes" if subject.test_id in handoff_ids else "No",
                review_imported=review.get("_review_present", "No"),
                ready_to_send=review.get("Ready To Send", ""),
                validation_status=review.get("Validation Status", ""),
                scheduled_send_at=scheduled_send_at,
                last_send_status=log.get("status", ""),
                last_send_time=log.get("timestamp", ""),
                template_type=subject.template_type,
                language=subject.language,
                ai_risk_flags=review.get("AI Risk Flags", ""),
                prompt_version=review.get("Prompt Version", ""),
                issues=_issue_summary(
                    recipient_email=recipient,
                    enabled="Yes" if subject.enabled else "No",
                    disabled_reason=_disabled_reason(subject.template_type, subject.personalisation_fact),
                    handoff_generated="Yes" if subject.test_id in handoff_ids else "No",
                    review_imported=review.get("_review_present", "No"),
                    ready_to_send=review.get("Ready To Send", ""),
                    validation_status=review.get("Validation Status", ""),
                    scheduled_send_at=scheduled_send_at,
                    last_send_status=log.get("status", ""),
                    prompt_version=review.get("Prompt Version", ""),
                ),
                next_action=_next_action(
                    recipient_email=recipient,
                    enabled="Yes" if subject.enabled else "No",
                    disabled_reason=_disabled_reason(subject.template_type, subject.personalisation_fact),
                    handoff_generated="Yes" if subject.test_id in handoff_ids else "No",
                    review_imported=review.get("_review_present", "No"),
                    ready_to_send=review.get("Ready To Send", ""),
                    validation_status=review.get("Validation Status", ""),
                    scheduled_send_at=scheduled_send_at,
                    last_send_status=log.get("status", ""),
                    prompt_version=review.get("Prompt Version", ""),
                ),
            )
        )
    return statuses


def load_file_lineage_status() -> FileLineageStatus:
    status, message = lineage_status(
        config.CHATGPT_HANDOFF_WORKBOOK_PATH,
        config.LOCAL_REVIEW_WORKBOOK_PATH,
    )
    return FileLineageStatus(status=status, message=message)


def _read_review_rows(review_workbook_path: Path) -> dict[str, dict[str, str]]:
    if not review_workbook_path.exists():
        return {}

    from openpyxl import load_workbook

    workbook = load_workbook(review_workbook_path, data_only=True)
    try:
        sheet = (
            workbook[LOCAL_REVIEW_SHEET_NAME]
            if LOCAL_REVIEW_SHEET_NAME in workbook.sheetnames
            else workbook.active
        )
        headers = {
            str(sheet.cell(1, column).value).strip(): column
            for column in range(1, sheet.max_column + 1)
            if sheet.cell(1, column).value
        }
        if "Test ID" not in headers:
            return {}

        rows: dict[str, dict[str, str]] = {}
        for row_number in range(2, sheet.max_row + 1):
            test_id = _cell(sheet, row_number, headers["Test ID"])
            if not test_id:
                continue
            rows[test_id] = {
                header: _cell(sheet, row_number, column)
                for header, column in headers.items()
            }
            rows[test_id]["_review_present"] = "Yes"
        return rows
    finally:
        workbook.close()


def _read_latest_log_status(send_log_path: Path) -> dict[tuple[str, str], dict[str, str]]:
    if not send_log_path.exists():
        return {}

    latest: dict[tuple[str, str], dict[str, str]] = {}
    with send_log_path.open("r", encoding="utf-8", newline="") as handle:
        for row in csv.DictReader(handle):
            test_id = row.get("test_id", "")
            recipient = row.get("recipient_email", "").lower()
            if test_id and recipient:
                if row.get("status", "") in TRANSIENT_STATUSES:
                    continue
                row["scheduled_send_at"] = clean_scheduled_send_at(row.get("scheduled_send_at", ""))
                if not row.get("scheduled_send_at", ""):
                    row["scheduled_send_at"] = scheduled_send_at_from_response(
                        row.get("graph_response", "")
                    )
                latest[(test_id, recipient)] = row
    return latest


def _issue_summary(
    *,
    recipient_email: str,
    enabled: str,
    disabled_reason: str = "",
    handoff_generated: str,
    review_imported: str,
    ready_to_send: str,
    validation_status: str,
    scheduled_send_at: str,
    last_send_status: str,
    prompt_version: str = "",
) -> str:
    issues: list[str] = []
    if not recipient_email.strip():
        issues.append("Missing email")
    if enabled != "Yes":
        if disabled_reason:
            issues.append(f"Disabled - {disabled_reason}")
        else:
            issues.append("Disabled")
    if last_send_status == "SENT":
        return ", ".join(issues)
    if handoff_generated != "Yes":
        issues.append("No handoff")
    if review_imported != "Yes":
        issues.append("No review")
    if ready_to_send != "Yes":
        issues.append("Not ready")
    if validation_status and validation_status != "VALID":
        issues.append("Invalid")
    if prompt_version and prompt_version != config.PROMPT_PACKET_VERSION:
        issues.append("Stale prompt")
    if last_send_status == "DRAFT_CREATED" and not scheduled_send_at.strip():
        issues.append("No schedule")
    if last_send_status == "SEND_STATUS_UNKNOWN":
        issues.append("Unknown send status")
    if last_send_status == "OUTLOOK_OUTBOX_LOCKED":
        issues.append("Outbox locked")
    if last_send_status == "DELIVERY_FAILED":
        issues.append("Delivery failed")
    if last_send_status == "REFUSED_TO_ANSWER":
        issues.append("Refused to answer")
    return ", ".join(issues)


def _disabled_reason(template_type: str, personalisation_fact: str) -> str:
    if template_type in config.DISABLED_TEMPLATE_REASONS:
        return config.DISABLED_TEMPLATE_REASONS[template_type]
    lowered_fact = personalisation_fact.lower()
    for keyword, reason in config.DISABLED_PERSONALISATION_KEYWORDS.items():
        if keyword in lowered_fact:
            return reason
    return ""


def _next_action(
    *,
    recipient_email: str,
    enabled: str,
    disabled_reason: str = "",
    handoff_generated: str,
    review_imported: str,
    ready_to_send: str,
    validation_status: str,
    scheduled_send_at: str,
    last_send_status: str,
    prompt_version: str = "",
) -> str:
    if not recipient_email.strip():
        return "Add recipient email"
    if enabled != "Yes":
        if disabled_reason:
            return "Keep disabled"
        return "Enable in source"
    if last_send_status == "SENT":
        return "Done"
    if last_send_status == "OUTLOOK_OUTBOX_LOCKED":
        return "Fix outbox send now"
    if last_send_status == "OUTLOOK_SCHEDULED":
        return "Check scheduled sent"
    if last_send_status == "DELIVERY_FAILED":
        return "Fix recipient email"
    if last_send_status == "REFUSED_TO_ANSWER":
        return "Keep excluded"
    if last_send_status == "SEND_STATUS_UNKNOWN":
        return "Manual send review"
    if handoff_generated != "Yes":
        return "Generate handoff"
    if review_imported != "Yes":
        return "Import review"
    if ready_to_send != "Yes":
        return "Review and mark ready"
    if validation_status != "VALID":
        return "Fix validation"
    if prompt_version and prompt_version != config.PROMPT_PACKET_VERSION:
        return "Regenerate handoff"
    if last_send_status == "DRAFT_CREATED":
        return "Set schedule" if not scheduled_send_at.strip() else "Schedule in Outlook"
    return "Make draft"


def _cell(sheet, row: int, column: int) -> str:
    value = sheet.cell(row, column).value
    if value is None:
        return ""
    return str(value).strip()

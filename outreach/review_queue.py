from __future__ import annotations

from datetime import datetime
from pathlib import Path

import config
from openpyxl import load_workbook

from outreach.logger import clean_scheduled_send_at
from outreach.models import QueueEmail
from outreach.review_importer import LOCAL_REVIEW_SHEET_NAME
from outreach.workbook_io import save_workbook_atomically


REQUIRED_COLUMNS = [
    "Test ID",
    "Target Name",
    "Recipient Email",
    "Recipient Type",
    "Template Type",
    "Language",
    "Subject",
    "Email Body",
    "Ready To Send",
    "Validation Status",
]


def read_ready_queue(review_workbook_path: Path) -> list[QueueEmail]:
    if not review_workbook_path.exists():
        raise FileNotFoundError(f"Review queue workbook not found: {review_workbook_path}")

    workbook = load_workbook(review_workbook_path, data_only=True)
    try:
        sheet, headers = _sheet_and_headers(workbook)
        _require_columns(headers, REQUIRED_COLUMNS)

        ready: list[QueueEmail] = []
        for row in range(2, sheet.max_row + 1):
            test_id = _cell(sheet, row, headers["Test ID"])
            if not test_id:
                continue

            ready_to_send = _is_yes(_cell(sheet, row, headers["Ready To Send"]))
            validation_status = _cell(sheet, row, headers["Validation Status"])
            scheduled_send_at = clean_scheduled_send_at(
                _optional_cell(sheet, row, headers, config.SCHEDULED_SEND_AT_COLUMN)
            )
            if not ready_to_send or validation_status != "VALID" or not is_due(scheduled_send_at):
                continue

            ready.append(_queue_email_from_row(sheet, row, headers, ready_to_send, scheduled_send_at))

        return ready
    finally:
        workbook.close()


def read_review_email(review_workbook_path: Path, test_id: str) -> QueueEmail:
    if not review_workbook_path.exists():
        raise FileNotFoundError(f"Review queue workbook not found: {review_workbook_path}")

    workbook = load_workbook(review_workbook_path, data_only=True)
    try:
        sheet, headers = _sheet_and_headers(workbook)
        _require_columns(headers, REQUIRED_COLUMNS)

        for row in range(2, sheet.max_row + 1):
            if _cell(sheet, row, headers["Test ID"]) == test_id:
                scheduled_send_at = clean_scheduled_send_at(
                    _optional_cell(sheet, row, headers, config.SCHEDULED_SEND_AT_COLUMN)
                )
                return _queue_email_from_row(
                    sheet,
                    row,
                    headers,
                    _is_yes(_cell(sheet, row, headers["Ready To Send"])),
                    scheduled_send_at,
                )
        raise ValueError(f"Test ID not found in review queue: {test_id}")
    finally:
        workbook.close()


def set_ready_to_send(review_workbook_path: Path, test_id: str, ready: bool) -> None:
    set_ready_to_send_many(review_workbook_path, {test_id: ready})


def set_ready_to_send_many(review_workbook_path: Path, ready_by_test_id: dict[str, bool]) -> None:
    if not review_workbook_path.exists():
        raise FileNotFoundError(f"Review queue workbook not found: {review_workbook_path}")
    if not ready_by_test_id:
        return

    workbook = load_workbook(review_workbook_path)
    try:
        sheet, headers = _sheet_and_headers(workbook)
        for header in ("Test ID", "Ready To Send"):
            if header not in headers:
                raise ValueError(f"Review queue missing column: {header}")

        remaining = set(ready_by_test_id)
        for row in range(2, sheet.max_row + 1):
            test_id = _cell(sheet, row, headers["Test ID"])
            if test_id in ready_by_test_id:
                sheet.cell(
                    row,
                    headers["Ready To Send"],
                    "Yes" if ready_by_test_id[test_id] else "No",
                )
                remaining.discard(test_id)
        if remaining:
            missing = ", ".join(sorted(remaining))
            raise ValueError(f"Test ID not found in review queue: {missing}")
        save_workbook_atomically(workbook, review_workbook_path)
    finally:
        workbook.close()


def set_scheduled_send_at(review_workbook_path: Path, test_id: str, scheduled_send_at: str) -> None:
    set_scheduled_send_at_many(review_workbook_path, {test_id: scheduled_send_at})


def set_scheduled_send_at_many(review_workbook_path: Path, scheduled_by_test_id: dict[str, str]) -> None:
    if not review_workbook_path.exists():
        raise FileNotFoundError(f"Review queue workbook not found: {review_workbook_path}")
    if not scheduled_by_test_id:
        return

    for scheduled_send_at in scheduled_by_test_id.values():
        if scheduled_send_at:
            parse_scheduled_send_at(scheduled_send_at)

    workbook = load_workbook(review_workbook_path)
    try:
        sheet, headers = _sheet_and_headers(workbook)
        if "Test ID" not in headers:
            raise ValueError("Review queue missing column: Test ID")
        schedule_column = _ensure_column(sheet, headers, config.SCHEDULED_SEND_AT_COLUMN)

        remaining = set(scheduled_by_test_id)
        for row in range(2, sheet.max_row + 1):
            test_id = _cell(sheet, row, headers["Test ID"])
            if test_id in scheduled_by_test_id:
                sheet.cell(row, schedule_column, scheduled_by_test_id[test_id])
                remaining.discard(test_id)
        if remaining:
            missing = ", ".join(sorted(remaining))
            raise ValueError(f"Test ID not found in review queue: {missing}")
        save_workbook_atomically(workbook, review_workbook_path)
    finally:
        workbook.close()


def remove_review_email(review_workbook_path: Path, test_id: str) -> bool:
    if not review_workbook_path.exists():
        return False

    workbook = load_workbook(review_workbook_path)
    try:
        sheet, headers = _sheet_and_headers(workbook)
        if "Test ID" not in headers:
            return False

        for row in range(2, sheet.max_row + 1):
            if _cell(sheet, row, headers["Test ID"]) == test_id:
                sheet.delete_rows(row, 1)
                save_workbook_atomically(workbook, review_workbook_path)
                return True
        return False
    finally:
        workbook.close()


def parse_scheduled_send_at(value: str) -> datetime:
    cleaned = clean_scheduled_send_at(value).replace("T", " ")
    try:
        return datetime.strptime(cleaned, "%Y-%m-%d %H:%M")
    except ValueError as exc:
        raise ValueError("Scheduled send time must use format YYYY-MM-DD HH:MM") from exc


def is_due(scheduled_send_at: str) -> bool:
    scheduled_send_at = clean_scheduled_send_at(scheduled_send_at)
    if not scheduled_send_at:
        return True
    return parse_scheduled_send_at(scheduled_send_at) <= datetime.now()


def _queue_email_from_row(sheet, row: int, headers: dict[str, int], ready_to_send: bool, scheduled_send_at: str) -> QueueEmail:
    return QueueEmail(
        test_id=_cell(sheet, row, headers["Test ID"]),
        target_name=_cell(sheet, row, headers["Target Name"]),
        recipient_email=_cell(sheet, row, headers["Recipient Email"]).lower(),
        recipient_type=_cell(sheet, row, headers["Recipient Type"]),
        template_type=_cell(sheet, row, headers["Template Type"]),
        language=_cell(sheet, row, headers["Language"]),
        subject=_cell(sheet, row, headers["Subject"]),
        body=_cell(sheet, row, headers["Email Body"]),
        ready_to_send=ready_to_send,
        validation_status=_cell(sheet, row, headers["Validation Status"]),
        scheduled_send_at=scheduled_send_at,
        ai_risk_flags=_optional_cell(sheet, row, headers, "AI Risk Flags"),
        prompt_version=_optional_cell(sheet, row, headers, "Prompt Version"),
    )


def _sheet_and_headers(workbook):
    sheet = (
        workbook[LOCAL_REVIEW_SHEET_NAME]
        if LOCAL_REVIEW_SHEET_NAME in workbook.sheetnames
        else workbook.active
    )
    headers = {
        str(sheet.cell(1, column).value).strip(): column
        for column in range(1, sheet.max_column + 1)
        if sheet.cell(1, column).value
    }
    return sheet, headers


def _require_columns(headers: dict[str, int], required: list[str]) -> None:
    missing = [header for header in required if header not in headers]
    if missing:
        raise ValueError(f"Review queue missing column(s): {', '.join(missing)}")


def _optional_cell(sheet, row: int, headers: dict[str, int], header: str) -> str:
    if header not in headers:
        return ""
    return _cell(sheet, row, headers[header])


def _ensure_column(sheet, headers: dict[str, int], header: str) -> int:
    if header in headers:
        return headers[header]
    column = sheet.max_column + 1
    sheet.cell(1, column, header)
    headers[header] = column
    return column


def _cell(sheet, row: int, column: int) -> str:
    value = sheet.cell(row, column).value
    if value is None:
        return ""
    return str(value).strip()


def _is_yes(value: str) -> bool:
    return value.strip().lower() in {"yes", "y", "true", "1"}
